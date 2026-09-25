from decimal import Decimal
from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.catalog.models import Product, ProductPart
from apps.common import packing_sheet as sheet_kit
from apps.masters.models import Merchant

from .models import Order, OrderLine


def under(sheet, caption):
    """What a heading box holds: the cell beneath its caption."""
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value == caption:
                return sheet.cell(row=cell.row + 1, column=cell.column).value
    raise AssertionError(f"no {caption!r} on the sheet")


class OrderFixture(APITestCase):
    """One order: six chairs two-to-a-box, and two tables of two parts each."""

    def setUp(self):
        self.merchant = Merchant.objects.create(code="UO", name="Urban Outfitters")

        self.chair = Product.objects.create(
            style_no="CHR-01", description="Chair", pack_per_box=2,
            box_length_in=40, box_width_in=30, box_height_in=24,
            product_weight_kg=Decimal("6"),
            box_weight_kg=Decimal("0.5"),
            packing_material_weight_kg=Decimal("0.2"),
        )
        self.table = Product.objects.create(
            style_no="TBL-01", description="Table", is_multi_part=True,
        )
        for name in ("Top", "Legs"):
            ProductPart.objects.create(
                product=self.table, name=name,
                box_length_in=40, box_width_in=30, box_height_in=24,
                product_weight_kg=Decimal("10"),
                box_weight_kg=Decimal("1"),
                packing_material_weight_kg=Decimal("0.5"),
            )

        self.order = Order.objects.create(name="Test", merchant=self.merchant)
        OrderLine.objects.create(
            order=self.order, product=self.chair, quantity=6, color="Charcoal Wash"
        )
        OrderLine.objects.create(
            order=self.order, product=self.table, quantity=2, color="Natural Oak"
        )

    def url(self, name):
        return reverse(f"order-{name}", args=[self.order.pk])

    def payload(self, cartons):
        """Reshape a GET response the way the packing screen posts it back."""
        return {
            "cartons": [
                {
                    "carton_no": carton["carton_no"],
                    "length_in": carton["length_in"],
                    "width_in": carton["width_in"],
                    "height_in": carton["height_in"],
                    "gross_weight_kg": carton["gross_weight_kg"],
                    "sort_order": index,
                    "contents": [
                        {
                            "product": item["product"],
                            "part": item["part"],
                            "description": item["description"],
                            "quantity": item["quantity"],
                            "unit": item["unit"],
                            "net_weight_kg": item["net_weight_kg"],
                        }
                        for item in carton["contents"]
                    ],
                }
                for index, carton in enumerate(cartons)
            ]
        }


class PackingApiTests(OrderFixture):
    def test_auto_pack_splits_by_parts_and_pack_per_box(self):
        response = self.client.post(self.url("auto-pack"))
        self.assertEqual(response.status_code, 200)
        # 6 chairs at 2/box = 3, plus 2 tables x 2 parts = 4.
        self.assertEqual(response.data["carton_count"], 7)
        self.assertTrue(response.data["can_save"])

    def test_a_parts_cartons_are_numbered_together(self):
        """One unbroken run per part, so the packing list can name a range
        rather than every second number."""
        self.client.post(self.url("auto-pack"))
        cartons = self.client.get(self.url("packing")).data["cartons"]

        by_part = {}
        for carton in cartons:
            for content in carton["contents"]:
                by_part.setdefault(content["description"], []).append(
                    carton["carton_no"]
                )

        self.assertEqual(by_part["Table — Top"], ["BOX-004", "BOX-005"])
        self.assertEqual(by_part["Table — Legs"], ["BOX-006", "BOX-007"])

    def test_packing_accepts_both_get_and_put(self):
        """Regression: a second @action on the same url_path shadowed the PUT."""
        self.client.post(self.url("auto-pack"))

        read = self.client.get(self.url("packing"))
        self.assertEqual(read.status_code, 200)

        saved = self.client.put(
            self.url("packing"), self.payload(read.data["cartons"]), format="json"
        )
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.data["carton_count"], 7)

    def test_saving_a_draft_plan_advances_the_order(self):
        self.client.post(self.url("auto-pack"))
        read = self.client.get(self.url("packing"))
        self.client.put(self.url("packing"), self.payload(read.data["cartons"]), format="json")

        self.order.refresh_from_db()
        self.assertEqual(self.order.status, "packing")

    def test_short_plan_is_rejected_and_rolled_back(self):
        self.client.post(self.url("auto-pack"))
        read = self.client.get(self.url("packing"))

        short = self.payload(read.data["cartons"][:-1])
        response = self.client.put(self.url("packing"), short, format="json")

        self.assertEqual(response.status_code, 400)
        self.assertTrue(
            any(b["code"] == "quantity_mismatch" for b in response.data["blockers"])
        )
        self.assertEqual(self.order.cartons.count(), 7)

    def test_multi_part_reconciliation_takes_the_minimum(self):
        """Three tops and two leg sets is two tables, not two and a half."""
        self.client.post(self.url("auto-pack"))
        top = self.table.parts.first()
        self.order.cartons.filter(contents__part=top).first().delete()

        response = self.client.get(self.url("packing"))
        row = next(
            r for r in response.data["reconciliation"] if r["style_no"] == "TBL-01"
        )
        self.assertEqual(row["packed"], 1)
        self.assertFalse(row["is_matched"])

    def test_a_carton_without_dimensions_is_blocked(self):
        """Regression: blank dimensions used to save cleanly and declare 0 CBM."""
        self.client.post(self.url("auto-pack"))
        read = self.client.get(self.url("packing"))

        payload = self.payload(read.data["cartons"])
        for carton in payload["cartons"]:
            carton["length_in"] = carton["width_in"] = carton["height_in"] = None

        response = self.client.put(self.url("packing"), payload, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertTrue(
            any(b["code"] == "missing_dimensions" for b in response.data["blockers"])
        )

    def test_cannot_mark_packed_while_blocked(self):
        response = self.client.post(self.url("advance-status"))
        self.assertEqual(response.status_code, 200)  # draft -> packing

        response = self.client.post(self.url("advance-status"))
        self.assertEqual(response.status_code, 400)  # packing -> packed is gated


class PackingListTests(OrderFixture):
    """The Excel document the shipping desk sends out."""

    # Column positions, as the sheet lays them out.
    RANGE, COUNT, STYLE, DESCRIPTION = 0, 1, 2, 3
    QTY, UNITS = 4, 5
    NNW, NET, GROSS = 6, 7, 8
    LENGTH, WIDTH, HEIGHT, CBM = 9, 10, 11, 12

    def rows(self):
        """Data rows only — between the column headers and the totals."""
        response = self.client.get(self.url("packing-list"))
        self.assertEqual(response.status_code, 200)

        sheet = load_workbook(BytesIO(b"".join(response.streaming_content))).active
        values = list(sheet.iter_rows(values_only=True))
        start = next(i for i, row in enumerate(values) if row[0] == "Carton Nos") + 1

        found = []
        for row in values[start:]:
            if row[0] and str(row[0]).startswith("TOTAL"):
                break
            found.append(row)
        return found

    def totals(self):
        """The footer row — the shipment, not a sum of the columns above."""
        response = self.client.get(self.url("packing-list"))
        sheet = load_workbook(BytesIO(b"".join(response.streaming_content))).active
        values = list(sheet.iter_rows(values_only=True))
        return next(
            row for row in values if row[0] and str(row[0]).startswith("TOTAL")
        )

    def repack(self, edit):
        """Auto-pack, let `edit` change the plan, then save it back."""
        self.client.post(self.url("auto-pack"))
        cartons = self.client.get(self.url("packing")).data["cartons"]
        edit(cartons)
        response = self.client.put(self.url("packing"), self.payload(cartons), format="json")
        self.assertEqual(response.status_code, 200, response.data)

    def test_it_needs_cartons_before_there_is_anything_to_list(self):
        response = self.client.get(self.url("packing-list"))
        self.assertEqual(response.status_code, 400)

    def test_identical_cartons_collapse_to_one_row(self):
        self.client.post(self.url("auto-pack"))
        rows = self.rows()

        # Seven cartons, but only three different things in them: the chairs,
        # the table tops and the leg sets.
        self.assertEqual(self.order.cartons.count(), 7)
        self.assertEqual(len(rows), 3)

        self.assertEqual(
            [row[self.DESCRIPTION] for row in rows],
            ["Chair", "Table — Top", "Table — Legs"],
        )

    def test_a_row_names_its_carton_range_and_counts_them(self):
        self.client.post(self.url("auto-pack"))
        chairs = self.rows()[0]

        self.assertEqual(chairs[self.RANGE], "BOX-001 – BOX-003")
        self.assertEqual(chairs[self.COUNT], 3)

    def test_each_part_gets_one_unbroken_range(self):
        self.order.lines.filter(product=self.table).update(quantity=4)
        self.client.post(self.url("auto-pack"))
        tops, legs = self.rows()[1], self.rows()[2]

        self.assertEqual(tops[self.RANGE], "BOX-004 – BOX-007")
        self.assertEqual(legs[self.RANGE], "BOX-008 – BOX-011")

    def test_a_hand_edited_plan_with_gaps_names_every_run(self):
        """Nothing is abbreviated away: each number a row covers is either
        printed or inside a printed run."""
        def renumber(cartons):
            # Push the third chair carton clear of the first two.
            cartons[2]["carton_no"] = "BOX-009"

        self.repack(renumber)
        chairs = self.rows()[0]

        self.assertEqual(chairs[self.RANGE], "BOX-001 – BOX-002, BOX-009")
        self.assertEqual(chairs[self.COUNT], 3)

    def test_a_row_describes_one_box_however_many_it_stands_for(self):
        self.client.post(self.url("auto-pack"))
        chairs = self.rows()[0]
        carton = self.order.cartons.first()

        self.assertEqual(chairs[self.COUNT], 3)
        self.assertEqual(chairs[self.QTY], 2)  # not 6
        self.assertEqual(chairs[self.UNITS], "PCS")
        self.assertEqual(Decimal(str(chairs[self.GROSS])), carton.gross_weight_kg)
        self.assertEqual(Decimal(str(chairs[self.NET])), carton.net_weight_kg)

    def test_the_footer_counts_the_boxes_the_columns_do_not(self):
        """
        Every figure on a row is one box's, so the shipment total cannot be
        the sum of the column — it has to multiply each row by its run.
        """
        self.client.post(self.url("auto-pack"))
        totals = self.totals()
        rows = self.rows()

        self.assertEqual(totals[self.COUNT], 7)
        self.assertEqual(
            totals[self.QTY],
            sum(row[self.QTY] * row[self.COUNT] for row in rows),
        )
        self.assertGreater(totals[self.QTY], sum(row[self.QTY] for row in rows))

    def test_a_row_carries_the_box_size_in_centimetres(self):
        self.client.post(self.url("auto-pack"))
        row = self.rows()[0]
        carton = self.order.cartons.first()

        # Excel hands the numbers back as floats, so compare on value.
        self.assertEqual(
            [Decimal(str(row[side])) for side in (self.LENGTH, self.WIDTH, self.HEIGHT)],
            [
                sheet_kit.cm(carton.length_in),
                sheet_kit.cm(carton.width_in),
                sheet_kit.cm(carton.height_in),
            ],
        )

    def test_the_printed_cbm_multiplies_out_from_the_printed_sides(self):
        """A broker rechecking the arithmetic on the page has to arrive at it."""
        self.client.post(self.url("auto-pack"))
        row = self.rows()[0]

        sides = Decimal(str(row[self.LENGTH]))
        sides *= Decimal(str(row[self.WIDTH]))
        sides *= Decimal(str(row[self.HEIGHT]))

        self.assertEqual(
            Decimal(str(row[self.CBM])),
            (sides / Decimal("1000000")).quantize(Decimal("0.0001")),
        )

    def test_a_carton_weighed_differently_gets_its_own_row(self):
        """The whole point of collapsing is that the merged cartons really
        are identical — one reweighed box has to break out."""
        def heavier(cartons):
            cartons[0]["gross_weight_kg"] = "13.400"

        self.repack(heavier)
        rows = self.rows()

        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0][self.RANGE], "BOX-001")
        self.assertEqual(rows[0][self.COUNT], 1)
        self.assertEqual(rows[1][self.RANGE], "BOX-002 – BOX-003")
        self.assertEqual(rows[1][self.COUNT], 2)


    def test_the_json_document_carries_exactly_what_the_sheet_does(self):
        """
        The print page and the workbook come off one builder, so a PDF signed
        at the desk cannot quote a figure the emailed sheet disagrees with.
        """
        self.client.post(self.url("auto-pack"))

        response = self.client.get(self.url("packing-list-data"))
        self.assertEqual(response.status_code, 200)
        document = response.data

        # openpyxl reads a written "" back as an empty cell.
        printed = [
            [value if value != "" else None for value in row["values"]]
            for block in document["blocks"]
            for row in block["rows"]
        ]
        self.assertEqual(printed, [list(row) for row in self.rows()])

        footer = list(self.totals())
        for column in (self.QTY, self.NNW, self.NET, self.GROSS, self.CBM):
            self.assertEqual(document["total"][column], footer[column])

    def test_the_json_document_needs_cartons_too(self):
        response = self.client.get(self.url("packing-list-data"))
        self.assertEqual(response.status_code, 400)

    def test_the_filename_names_the_order(self):
        self.client.post(self.url("auto-pack"))
        response = self.client.get(self.url("packing-list"))
        self.assertIn(
            f"packing-list-{self.order.number}.xlsx", response["Content-Disposition"]
        )


class InvoiceCodeTests(OrderFixture):
    """The buyer's customs clear on the US tariff code, not India's HSN."""

    def test_the_invoice_prints_the_hts_code_not_the_hsn(self):
        Product.objects.filter(pk=self.chair.pk).update(
            hsn_code="9401", hts_code="9401.61.6011"
        )
        # A multi-part product has no code of its own; its parts carry it.
        self.table.parts.update(hsn_code="9403", hts_code="9403.60.8081")
        self.client.post(self.url("auto-pack"))

        response = self.client.get(self.url("invoice-data"))
        self.assertEqual(response.status_code, 200)

        codes = {line[1]: line[2] for line in response.data["lines"]}
        self.assertEqual(codes, {"CHR-01": "9401.61.6011", "TBL-01": "9403.60.8081"})


class ExportHeadingTests(OrderFixture):
    """Hardgoods has no form for export details, so the order's address stands in."""

    def sheet(self, name):
        response = self.client.get(self.url(name))
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(b"".join(response.streaming_content))).active

    def test_the_ship_to_address_fills_other_consignee(self):
        Order.objects.filter(pk=self.order.pk).update(
            ship_line1="190 Yarnell Road", ship_city="Pottstown",
            ship_state="PA", ship_postal_code="19465",
        )
        self.client.post(self.url("auto-pack"))

        for name in ("invoice", "packing-list"):
            self.assertEqual(
                under(self.sheet(name), "Other Consignee (Shipp To-)"),
                "190 Yarnell Road\nPottstown, PA 19465\nUS",
            )

    def test_an_other_consignee_typed_for_the_shipment_wins(self):
        Order.objects.filter(pk=self.order.pk).update(
            ship_line1="190 Yarnell Road",
            export_details={"other_consignee": "L&J Transportation"},
        )
        self.client.post(self.url("auto-pack"))

        self.assertEqual(
            under(self.sheet("invoice"), "Other Consignee (Shipp To-)"),
            "L&J Transportation",
        )
