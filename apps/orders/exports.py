"""
The packing list as the shipping desk sends it: one row per distinct thing
packed, not per carton. Twelve identical cartons of one chair collapse to a
single row carrying their carton range and count, so the sheet is as long as
the order has different items rather than as long as it has boxes.

The grouping and the sheet furniture live in `apps/common/packing_sheet.py`,
shared with Display.
"""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook

from apps.common import invoice_sheet as invoice_kit
from apps.common import packing_sheet as sheet_kit

from . import services
from .models import Order


def invoice_filename(order: Order) -> str:
    return f"invoice-{order.number}.xlsx"


def packing_list_filename(order: Order) -> str:
    return f"packing-list-{order.number}.xlsx"


def packing_document(order: Order) -> sheet_kit.Document:
    cartons = list(
        order.cartons.prefetch_related("contents__product", "contents__part")
    )

    totals = sheet_kit.Totals()
    rows = []
    for group in sheet_kit.group_cartons(cartons):
        totals.add(group)
        for position, values in enumerate(group.rows()):
            rows.append((values, position == 0))

    count = totals.cartons
    return sheet_kit.Document(
        title="PACKING LIST",
        order_title=order.name,
        facts=[
            ("Order", order.number),
            ("Name", order.name),
            ("Merchant", order.merchant.name),
            ("Buyer", order.buyer_name),
            ("Ship to", _address(order)),
            ("Status", order.get_status_display()),
        ],
        blocks=[sheet_kit.Block(rows=rows)],
        total=totals,
        total_label=f"TOTAL · all {count} box{'es' if count != 1 else ''}",
        export_details=order.export_details or {},
        ship_to=_ship_to(order),
    )


def build_packing_list(order: Order) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = order.number

    sheet_kit.write_document(sheet, packing_document(order))

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _address(order: Order) -> str:
    parts = [
        order.ship_line1,
        order.ship_line2,
        order.ship_city,
        order.ship_state,
        order.ship_postal_code,
        order.ship_country,
    ]
    return ", ".join(part for part in parts if part)


def _ship_to(order: Order) -> list[str]:
    """
    The order's own address, for the heading's Other Consignee when the
    export details name nobody. A country alone is no address to ship to.
    """
    if not order.ship_line1:
        return []
    town = " ".join(part for part in (order.ship_state, order.ship_postal_code) if part)
    town = ", ".join(part for part in (order.ship_city, town) if part)
    lines = (order.ship_line1, order.ship_line2, town, order.ship_country)
    return [line for line in lines if line]


def invoice_document(order: Order) -> invoice_kit.Document:
    """
    One row per style across the whole order, priced from the order lines.

    Billed on what is **packed**, not what was ordered, and counted the way
    `reconcile` counts it — a multi-part style is a unit only once every
    piece of it has a box. That keeps this document and the packing list
    agreeing on a single number, which is the pair a broker checks first.
    """
    lines = {line.product_id: line for line in order.lines.select_related("product")}
    invoice = invoice_kit.Invoice()

    for reconciled in services.reconcile(order):
        if reconciled.packed == 0:
            continue
        line = lines[reconciled.product]
        product = line.product
        invoice.add(
            product.id,
            style_no=product.style_no,
            hts_code=_hts(product),
            description=_customs(product),
            quantity=reconciled.packed,
            net_weight_kg=_unit_weight(product) * reconciled.packed,
            rate=line.rate_usd,
        )

    cartons = list(order.cartons.all())
    return invoice_kit.Document(
        invoice=invoice,
        order_title=order.name,
        export_details=order.export_details or {},
        ship_to=_ship_to(order),
        marks=sheet_kit.carton_range(cartons),
        boxes=len(cartons),
    )


def build_invoice(order: Order) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoice"

    invoice_kit.write_document(sheet, invoice_document(order))

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _hts(product) -> str:
    """
    The US tariff code, never the HSN: the buyer's customs clear on HTS. A
    part's code where the product has none — they may differ per part.
    """
    if product.hts_code:
        return product.hts_code
    return next((part.hts_code for part in product.parts.all() if part.hts_code), "")


def _customs(product) -> str:
    return (
        product.customs_description
        or product.style_name
        or product.description
        or product.style_no
    )


def _unit_weight(product) -> Decimal:
    """One saleable unit: the item, or every part of it together."""
    if product.is_multi_part:
        return sum(
            (part.product_weight_kg or Decimal("0") for part in product.parts.all()),
            start=Decimal("0"),
        )
    return product.product_weight_kg or Decimal("0")
