import base64
import tempfile
from collections import Counter
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from rest_framework.test import APITestCase

from apps.masters.models import Merchant, Store

from . import services
from .models import (
    DisplayCarton,
    DisplayCartonContent,
    DisplayOrder,
    DisplayOrderLine,
    DisplayProduct,
    DisplayProductImage,
    DisplayProductPart,
    PackTemplate,
    PackTemplateItem,
)


def product(style_no, weight, size):
    length, width, height = size
    return DisplayProduct.objects.create(
        style_no=style_no,
        description=style_no,
        product_weight_kg=Decimal(str(weight)),
        length_in=Decimal(str(length)),
        width_in=Decimal(str(width)),
        height_in=Decimal(str(height)),
    )


def display_part(**overrides):
    """A part complete enough to save. Tests override what they assert on."""
    return {
        "name": "Part",
        "customs_description": "Decorative display part",
        "hsn_code": "9505",
        "hts_code": "9505.10.5020",
        "product_weight_kg": "0.400",
        "length_in": "12.00",
        "width_in": "8.00",
        "height_in": "3.00",
        "sort_order": 0,
        **overrides,
    }


def display_product(**overrides):
    """A single-piece display product complete enough to save."""
    return {
        "style_no": "DSP-NEW-00",
        "description": "Display item",
        "customs_description": "Decorative display item",
        "hsn_code": "9505",
        "hts_code": "9505.10.5020",
        "product_weight_kg": "0.500",
        "length_in": "24.00",
        "width_in": "24.00",
        "height_in": "5.00",
        **overrides,
    }


def store(name="MAIN"):
    return Store.objects.create(name=name)


def template(code, items, box=(28, 20, 16), box_kg=0.8, packing_kg=0.25):
    length, width, height = box
    tpl = PackTemplate.objects.create(
        code=code,
        name=code,
        box_length_in=Decimal(str(length)),
        box_width_in=Decimal(str(width)),
        box_height_in=Decimal(str(height)),
        box_weight_kg=Decimal(str(box_kg)),
        packing_material_weight_kg=Decimal(str(packing_kg)),
        is_library=True,
    )
    for prod, quantity in items:
        PackTemplateItem.objects.create(template=tpl, product=prod, quantity=quantity)
    return tpl


def under(sheet, caption):
    """What a heading box holds: the cell beneath its caption."""
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value == caption:
                return sheet.cell(row=cell.row + 1, column=cell.column).value
    raise AssertionError(f"no {caption!r} on the sheet")


class DisplayPackingTests(TestCase):
    def setUp(self):
        self.merchant = Merchant.objects.create(code="TRN", name="Terrain Home")
        self.store = store()

        self.bow = product("DSP-BOW-12", 0.08, (12, 12, 4))
        self.orn = product("DSP-ORN-06", 0.55, (9, 6, 3))
        self.wrt = product("DSP-WRT-24", 1.20, (24, 24, 5))
        self.grl = product("DSP-GRL-72", 0.90, (14, 10, 8))

        self.order = DisplayOrder.objects.create(name="Holiday", merchant=self.merchant)
        for prod, quantity in (
            (self.bow, 840),
            (self.orn, 96),
            (self.wrt, 140),
            (self.grl, 158),
        ):
            DisplayOrderLine.objects.create(
                order=self.order, store=self.store, product=prod, quantity=quantity
            )

        self.t1 = template(
            "TPL-001",
            [(self.bow, 30), (self.wrt, 5)],
            box=(44, 32, 32),
            box_kg=1.6,
            packing_kg=0.5,
        )
        self.t2 = template(
            "TPL-002",
            [(self.orn, 12), (self.grl, 6)],
            box=(30, 20, 20),
            box_kg=0.9,
            packing_kg=0.6,
        )
        self.t3 = template("TPL-003", [(self.grl, 6)])
        self.t4 = template(
            "TPL-004", [(self.wrt, 10)], box=(42, 30, 30), box_kg=1.5, packing_kg=0.5
        )

    # ── The wall ─────────────────────────────────────────────────────

    def test_max_applications_takes_the_scarcest_product(self):
        """12 per box against 96 ornament sets allows 8; the garlands allow 26."""
        remaining = services.store_remaining(self.order, self.store.id)
        self.assertEqual(services.max_applications(self.t2, remaining), 8)

    def test_template_with_no_items_never_applies(self):
        empty = PackTemplate.objects.create(code="TPL-EMPTY", name="empty")
        remaining = services.store_remaining(self.order, self.store.id)
        self.assertEqual(services.max_applications(empty, remaining), 0)

    # ── The loop ─────────────────────────────────────────────────────

    def test_the_worked_example(self):
        """ARCHITECTURE.md §10.2, end to end."""
        for tpl, count in ((self.t1, 28), (self.t2, 8), (self.t3, 18)):
            step = services.apply_step(self.order, tpl, self.store)
            self.assertEqual(step.count, count, f"{tpl.code} applied {step.count}")

        self.assertEqual(services.remaining_quantities(self.order), {(self.store.id, self.grl.id, None): 2})

        tail = template("TPL-TAIL", [(self.grl, 2)], box=(20, 14, 12))
        services.apply_step(self.order, tail, self.store)

        self.assertEqual(services.remaining_quantities(self.order), {})
        self.assertEqual(self.order.cartons.count(), 55)
        self.assertEqual(services.find_blockers(self.order), [])
        self.assertTrue(all(row.is_matched for row in services.reconcile(self.order)))

    def test_a_count_above_capacity_is_refused(self):
        with self.assertRaises(services.PackingError):
            services.apply_step(self.order, self.t1, self.store, count=29)

    def test_a_template_that_does_not_fit_is_refused(self):
        services.apply_step(self.order, self.t1, self.store)  # 28, takes every wreath
        with self.assertRaises(services.PackingError):
            services.apply_step(self.order, self.t4, self.store)

    # ── Cartons ──────────────────────────────────────────────────────

    def test_carton_weights_come_from_contents(self):
        services.apply_step(self.order, self.t1, self.store, count=1)
        carton = self.order.cartons.get()

        # 30 x 0.08 + 5 x 1.20 = 8.4, plus 0.5 packing, plus 1.6 box.
        self.assertEqual(carton.net_weight_kg, Decimal("8.900"))
        self.assertEqual(carton.gross_weight_kg, Decimal("10.500"))
        self.assertEqual(carton.contents.count(), 2)

    def test_carton_numbers_run_on_across_steps(self):
        services.apply_step(self.order, self.t1, self.store, count=2)
        services.apply_step(self.order, self.t2, self.store, count=2)
        self.assertEqual(
            list(self.order.cartons.values_list("carton_no", flat=True)),
            ["BOX-001", "BOX-002", "BOX-003", "BOX-004"],
        )

    def test_earlier_carton_numbers_survive_a_later_edit(self):
        """A number already written on a box must not move."""
        services.apply_step(self.order, self.t1, self.store, count=2)
        services.apply_step(self.order, self.t2, self.store, count=2)
        services.recount_step(self.order, 2, 1)

        self.assertEqual(
            list(
                self.order.cartons.filter(step__sequence=1).values_list(
                    "carton_no", flat=True
                )
            ),
            ["BOX-001", "BOX-002"],
        )

    # ── Editing the plan ─────────────────────────────────────────────

    def test_recount_clamps_a_starved_later_step_and_reports_it(self):
        small = DisplayOrder.objects.create(name="Small", merchant=self.merchant)
        DisplayOrderLine.objects.create(
            order=small, store=self.store, product=self.bow, quantity=10
        )
        a = template("TPL-A", [(self.bow, 5)])
        b = template("TPL-B", [(self.bow, 5)])

        services.apply_step(small, a, self.store, count=1)
        services.apply_step(small, b, self.store, count=1)

        adjustments = services.recount_step(small, 1, 2)

        self.assertEqual(len(adjustments), 1)
        self.assertEqual((adjustments[0].template_code, adjustments[0].now), ("TPL-B", 0))
        self.assertEqual(small.steps.count(), 1)
        self.assertEqual(services.remaining_quantities(small), {})

    def test_deleting_a_step_closes_the_gap_in_the_sequence(self):
        services.apply_step(self.order, self.t1, self.store, count=1)
        services.apply_step(self.order, self.t2, self.store, count=1)
        services.apply_step(self.order, self.t3, self.store, count=1)

        services.delete_step(self.order, 2)

        self.assertEqual(
            list(self.order.steps.values_list("sequence", "template__code")),
            [(1, "TPL-001"), (2, "TPL-003")],
        )

    def test_deleting_a_step_returns_its_units_to_the_remainder(self):
        services.apply_step(self.order, self.t2, self.store, count=8)
        before = services.remaining_quantities(self.order)[(self.store.id, self.grl.id, None)]
        services.delete_step(self.order, 1)
        after = services.remaining_quantities(self.order)[(self.store.id, self.grl.id, None)]
        self.assertEqual(after - before, 48)

    # ── Reconciliation, blockers, warnings ───────────────────────────

    def test_an_unpacked_remainder_blocks(self):
        services.apply_step(self.order, self.t1, self.store, count=1)
        codes = {blocker.code for blocker in services.find_blockers(self.order)}
        self.assertIn("quantity_mismatch", codes)

    def test_a_carton_missing_dimensions_blocks(self):
        services.apply_step(self.order, self.t1, self.store, count=1)
        self.order.cartons.update(width_in=None)
        codes = {blocker.code for blocker in services.find_blockers(self.order)}
        self.assertIn("missing_dimensions", codes)

    def test_gross_below_net_blocks(self):
        services.apply_step(self.order, self.t1, self.store, count=1)
        self.order.cartons.update(gross_weight_kg=Decimal("1.000"))
        codes = {blocker.code for blocker in services.find_blockers(self.order)}
        self.assertIn("gross_below_net", codes)

    def test_a_mostly_empty_box_warns_but_does_not_block(self):
        roomy = template("TPL-AIR", [(self.bow, 1)], box=(40, 40, 40))
        services.apply_step(self.order, roomy, self.store, count=1)

        codes = {warning.code for warning in services.find_warnings(self.order)}
        self.assertIn("poor_fill", codes)
        self.assertNotIn(
            "poor_fill", {b.code for b in services.find_blockers(self.order)}
        )

    def test_overfilling_a_carton_warns_but_does_not_block(self):
        services.apply_step(self.order, self.t3, self.store, count=1)
        content = self.order.cartons.get().contents.get()
        content.quantity = 8  # template holds 6
        content.save()

        codes = {warning.code for warning in services.find_warnings(self.order)}
        self.assertIn("template_capacity_exceeded", codes)

    # ── The picker ───────────────────────────────────────────────────

    def test_only_templates_that_still_fit_are_offered(self):
        services.apply_step(self.order, self.t1, self.store)
        services.apply_step(self.order, self.t2, self.store)
        services.apply_step(self.order, self.t3, self.store)

        self.assertEqual(services.applicable_templates(self.order, self.store.id), [])

    def test_a_one_off_is_offered_on_the_order_it_was_written_for(self):
        """Saving a tail-filler that then never appears is the same as losing it."""
        tail = template("TPL-TAIL", [(self.grl, 1)])
        tail.is_library = False
        tail.order = self.order
        tail.save()

        offered = {row["code"] for row in services.applicable_templates(self.order, self.store.id)}
        self.assertIn("TPL-TAIL", offered)

    def test_a_one_off_stays_off_every_other_order(self):
        tail = template("TPL-TAIL", [(self.grl, 1)])
        tail.is_library = False
        tail.order = self.order
        tail.save()

        other = DisplayOrder.objects.create(name="Other", merchant=self.merchant)
        DisplayOrderLine.objects.create(
            order=other, store=self.store, product=self.grl, quantity=50
        )

        offered = {row["code"] for row in services.applicable_templates(other, self.store.id)}
        self.assertNotIn("TPL-TAIL", offered)

    def test_another_merchants_template_is_not_offered(self):
        other = Merchant.objects.create(code="WE", name="West Elm")
        mine = template("TPL-MINE", [(self.grl, 1)])
        mine.merchant = other
        mine.save()

        offered = {row["code"] for row in services.applicable_templates(self.order, self.store.id)}
        self.assertNotIn("TPL-MINE", offered)

    # ── Remaining follows the cartons, not the steps ─────────────────

    def test_a_hand_edited_carton_moves_the_remainder(self):
        services.apply_step(self.order, self.t1, self.store, count=1)
        before = services.remaining_quantities(self.order)[(self.store.id, self.bow.id, None)]

        carton = self.order.cartons.get()
        content = carton.contents.get(product=self.bow)
        content.quantity += 5
        content.save()

        after = services.remaining_quantities(self.order)[(self.store.id, self.bow.id, None)]
        self.assertEqual(before - after, 5)


class OrderLineStoreTests(APITestCase):
    """
    An order is split across a merchant's stores, and each store keeps its own
    mix — two stores need not want the same products at all.
    """

    def setUp(self):
        self.merchant = Merchant.objects.create(code="ANT", name="Anthropologie")
        self.portland = store("118")
        self.austin = store("204")
        self.wreath = product("DSP-WRT-24", 1.20, (24, 24, 5))
        self.tree = product("DSP-TRE-60", 3.40, (30, 22, 6))

    def payload(self, lines):
        return {
            "name": "Winter decor",
            "merchant": self.merchant.id,
            "buyer_name": "Anthropologie Buying",
            "lines": lines,
        }

    def line(self, store_obj, product_obj, quantity):
        return {
            "store": store_obj.id,
            "product": product_obj.id,
            "quantity": quantity,
        }

    def test_two_stores_may_take_different_products(self):
        response = self.client.post(
            "/api/display-orders/",
            self.payload(
                [
                    self.line(self.portland, self.wreath, 24),
                    self.line(self.austin, self.tree, 12),
                ]
            ),
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        rows = {(r["store_name"], r["product_style_no"]) for r in response.data["lines"]}
        self.assertEqual(rows, {("118", "DSP-WRT-24"), ("204", "DSP-TRE-60")})

    def test_the_same_product_may_repeat_across_stores(self):
        response = self.client.post(
            "/api/display-orders/",
            self.payload(
                [
                    self.line(self.portland, self.wreath, 24),
                    self.line(self.austin, self.wreath, 18),
                ]
            ),
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(len(response.data["lines"]), 2)

    def test_one_store_cannot_list_a_product_twice(self):
        response = self.client.post(
            "/api/display-orders/",
            self.payload(
                [
                    self.line(self.portland, self.wreath, 24),
                    self.line(self.portland, self.wreath, 6),
                ]
            ),
            format="json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("lines", response.data)

    def test_any_store_may_go_on_any_merchants_order(self):
        """Stores are global — the order's merchant does not narrow the list."""
        response = self.client.post(
            "/api/display-orders/",
            self.payload([self.line(store("900"), self.wreath, 24)]),
            format="json",
        )
        self.assertEqual(response.status_code, 201)

    def test_reconciliation_keeps_the_stores_apart(self):
        """
        Twenty-four wreaths in Portland do nothing for Austin, so the same
        product ordered by two stores is two rows, not one total.
        """
        order = DisplayOrder.objects.create(name="Split", merchant=self.merchant)
        for outlet, quantity in ((self.portland, 24), (self.austin, 18)):
            DisplayOrderLine.objects.create(
                order=order, store=outlet, product=self.wreath, quantity=quantity
            )

        rows = services.reconcile(order)
        self.assertEqual(
            {(row.store_name, row.ordered) for row in rows},
            {("118", 24), ("204", 18)},
        )

        scoped = services.reconcile(order, self.portland.id)
        self.assertEqual([row.ordered for row in scoped], [24])

    def test_a_store_with_order_lines_cannot_be_deleted(self):
        self.client.post(
            "/api/display-orders/",
            self.payload([self.line(self.portland, self.wreath, 24)]),
            format="json",
        )
        # PROTECT — the store is the address those cartons ship to.
        with self.assertRaises(Exception):
            self.portland.delete()

    def test_stores_keep_the_order_the_sheet_gave_them(self):
        """
        Neither by name ("1839" before "804") nor by number: the sheet's own
        column order, which is the first time the payload names each store.
        """
        first, second, third = store("900"), store("1839"), store("804")
        response = self.client.post(
            "/api/display-orders/",
            self.payload(
                [
                    self.line(first, self.wreath, 6),
                    self.line(second, self.wreath, 6),
                    self.line(first, self.tree, 2),
                    self.line(third, self.wreath, 6),
                ]
            ),
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)

        order = DisplayOrder.objects.get(pk=response.data["id"])
        names = [outlet.name for outlet in services.order_stores(order)]
        self.assertEqual(names, ["900", "1839", "804"])
        self.assertEqual(
            [row["store_name"] for row in response.data["lines"]],
            ["900", "900", "1839", "804"],
        )

    def test_an_edit_can_change_the_store_order(self):
        created = self.client.post(
            "/api/display-orders/",
            self.payload(
                [
                    self.line(self.portland, self.wreath, 6),
                    self.line(self.austin, self.wreath, 6),
                ]
            ),
            format="json",
        ).data

        response = self.client.patch(
            f"/api/display-orders/{created['id']}/",
            {
                "lines": [
                    self.line(self.austin, self.wreath, 6),
                    self.line(self.portland, self.wreath, 6),
                ]
            },
            format="json",
        )
        self.assertEqual(response.status_code, 200, response.data)

        order = DisplayOrder.objects.get(pk=created["id"])
        self.assertEqual(
            [outlet.name for outlet in services.order_stores(order)], ["204", "118"]
        )


class StorePackingTests(TestCase):
    """
    A carton never holds two stores' goods, so the greedy loop runs inside a
    store. Demand cannot be pooled and a step belongs to one destination.
    """

    def setUp(self):
        self.merchant = Merchant.objects.create(code="ANT", name="Anthropologie")
        self.portland = store("118")
        self.austin = store("204")
        self.bow = product("DSP-BOW-12", 0.08, (12, 12, 4))

        self.order = DisplayOrder.objects.create(name="Split", merchant=self.merchant)
        for outlet, quantity in ((self.portland, 30), (self.austin, 30)):
            DisplayOrderLine.objects.create(
                order=self.order, store=outlet, product=self.bow, quantity=quantity
            )

        self.box60 = template("TPL-60", [(self.bow, 60)])
        self.box30 = template("TPL-30", [(self.bow, 30)])

    def test_demand_is_never_pooled_across_stores(self):
        """
        Sixty bows are ordered but split thirty apiece, so a sixty-per-box
        design fits nobody — the pooled total is a number no carton can use.
        """
        self.assertEqual(
            services.max_applications(
                self.box60, services.store_remaining(self.order, self.portland.id)
            ),
            0,
        )
        with self.assertRaises(services.PackingError):
            services.apply_step(self.order, self.box60, self.portland)

    def test_a_design_is_only_offered_where_it_fits(self):
        offered = {
            row["code"]
            for row in services.applicable_templates(self.order, self.portland.id)
        }
        self.assertEqual(offered, {"TPL-30"})

    def test_packing_one_store_leaves_the_other_untouched(self):
        services.apply_step(self.order, self.box30, self.portland, count=1)

        remaining = services.remaining_quantities(self.order)
        self.assertNotIn((self.portland.id, self.bow.id, None), remaining)
        self.assertEqual(remaining[(self.austin.id, self.bow.id, None)], 30)

    def test_cartons_and_steps_carry_their_store(self):
        step = services.apply_step(self.order, self.box30, self.austin, count=1)

        self.assertEqual(step.store_id, self.austin.id)
        self.assertEqual(
            list(self.order.cartons.values_list("store_id", flat=True)),
            [self.austin.id],
        )

    def test_carton_numbers_run_on_across_stores(self):
        """One shipment, one run of numbers — the store is a label, not a reset."""
        services.apply_step(self.order, self.box30, self.portland, count=1)
        services.apply_step(self.order, self.box30, self.austin, count=1)

        self.assertEqual(
            list(self.order.cartons.values_list("carton_no", flat=True)),
            ["BOX-001", "BOX-002"],
        )

    def test_the_summary_reports_each_store_separately(self):
        services.apply_step(self.order, self.box30, self.portland, count=1)

        rows = {row["store_name"]: row for row in services.store_summaries(self.order)}
        self.assertTrue(rows["118"]["is_done"])
        self.assertEqual(rows["118"]["carton_count"], 1)
        self.assertFalse(rows["204"]["is_done"])
        self.assertEqual(rows["204"]["remaining"], 30)

    def test_a_plan_copies_onto_every_store_with_the_same_demand(self):
        """Portland and Austin both want thirty, so one plan serves both."""
        services.apply_step(self.order, self.box30, self.austin, count=1)

        copied = services.replicate_plan(self.order, self.austin.id)

        self.assertEqual([s.name for s in copied], ["118"])
        self.assertEqual(self.order.steps.filter(store=self.portland).count(), 1)
        self.assertEqual(services.remaining_quantities(self.order), {})

    def test_a_store_wanting_something_else_is_not_copied_onto(self):
        odd = store("331")
        DisplayOrderLine.objects.create(
            order=self.order, store=odd, product=self.bow, quantity=45
        )
        services.apply_step(self.order, self.box30, self.austin, count=1)

        targets = services.replicable_stores(self.order, self.austin.id)
        self.assertEqual([s.name for s in targets], ["118"])

    def test_a_part_packed_store_is_left_alone(self):
        """Copying onto a store somebody has already started would overpack it."""
        services.apply_step(self.order, self.box30, self.austin, count=1)
        services.apply_step(self.order, self.box30, self.portland, count=1)

        self.assertEqual(services.replicable_stores(self.order, self.austin.id), [])
        with self.assertRaises(services.PackingError):
            services.replicate_plan(self.order, self.austin.id)

    def test_a_plan_scoped_to_a_store_shows_only_its_steps(self):
        services.apply_step(self.order, self.box30, self.portland, count=1)
        services.apply_step(self.order, self.box30, self.austin, count=1)

        plan = services.packing_summary(self.order, self.portland.id)
        self.assertEqual([row.store_name for row in plan["steps"]], ["118"])
        self.assertEqual(len(services.packing_summary(self.order)["steps"]), 2)


class DisplayPackingListTests(APITestCase):
    """The sheet is blocked by store, because that is how it is picked."""

    def setUp(self):
        self.merchant = Merchant.objects.create(code="ANT", name="Anthropologie")
        self.portland = store("118")
        self.portland.ship_city = "Portland"
        self.portland.save()
        self.austin = store("204")
        self.bow = product("DSP-BOW-12", 0.08, (12, 12, 4))

        self.order = DisplayOrder.objects.create(name="Split", merchant=self.merchant)
        for outlet in (self.portland, self.austin):
            DisplayOrderLine.objects.create(
                order=self.order, store=outlet, product=self.bow, quantity=30
            )
        self.tpl = template("TPL-30", [(self.bow, 30)])

    def url(self):
        return f"/api/display-orders/{self.order.id}/packing-list/"

    def sheet(self, document="packing-list"):
        from openpyxl import load_workbook

        response = self.client.get(f"/api/display-orders/{self.order.id}/{document}/")
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(b"".join(response.streaming_content))).active

    def rows(self):
        from openpyxl import load_workbook

        response = self.client.get(self.url())
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(b"".join(response.streaming_content)))
        return [
            [cell.value for cell in row] for row in book[self.order.number].iter_rows()
        ]

    def table(self):
        """The header, the item rows under it, and the total row last."""
        rows = self.rows()
        at = next(index for index, row in enumerate(rows) if "Style No" in row)
        return rows[at], rows[at + 1 : -1], rows[-1]

    def column(self, label):
        header, body, _ = self.table()
        return [row[header.index(label)] for row in body]

    def test_the_columns_follow_the_hand_made_list(self):
        services.apply_step(self.order, self.tpl, self.portland, count=1)
        header, _, _ = self.table()
        self.assertEqual(
            header,
            [
                "Carton Nos", "Total No of Boxes", "Store No", "Style No",
                "Customs Description", "Qty / Box", "Units", "NNW (kg)",
                "N.W. (kg)", "G.W. (kg)", "L (in)", "W (in)", "H (in)",
                "L (cm)", "W (cm)", "H (cm)", "CBM",
            ],
        )

    def test_the_json_document_carries_what_the_sheet_does(self):
        """One builder feeds both, so the PDF and the .xlsx cannot drift."""
        services.apply_step(self.order, self.tpl, self.portland, count=1)
        services.apply_step(self.order, self.tpl, self.austin, count=1)

        response = self.client.get(
            f"/api/display-orders/{self.order.id}/packing-list-data/"
        )
        self.assertEqual(response.status_code, 200)
        document = response.data

        header, body, _ = self.table()
        self.assertEqual([column["label"] for column in document["columns"]], header)
        self.assertEqual(
            [row["values"] for block in document["blocks"] for row in block["rows"]],
            body,
        )

    def test_no_cartons_is_a_400_not_an_empty_sheet(self):
        response = self.client.get(self.url())
        self.assertEqual(response.status_code, 400)

    def test_the_store_is_named_once_where_its_run_starts(self):
        # Two cartons for one store, so the range notation is exercised too.
        self.order.lines.filter(store=self.portland).update(quantity=60)
        services.apply_step(self.order, self.tpl, self.portland, count=2)
        services.apply_step(self.order, self.tpl, self.austin, count=1)

        self.assertEqual(self.column("Store No"), [118, 204])
        self.assertEqual(
            self.column("Carton Nos"), ["1 – 2", 3]
        )
        self.assertEqual(self.column("Total No of Boxes"), [2, 1])

    def test_a_mixed_box_names_its_store_only_on_the_first_row(self):
        tree = product("DSP-TRE-60", 3.40, (30, 22, 6))
        DisplayOrderLine.objects.create(
            order=self.order, store=self.portland, product=tree, quantity=2
        )
        mixed = template("TPL-MIX", [(self.bow, 30), (tree, 2)])
        services.apply_step(self.order, mixed, self.portland, count=1)

        self.assertEqual(self.column("Store No"), [118, None])
        self.assertEqual(self.column("Carton Nos"), [1, None])

    def test_a_box_is_measured_in_inches_and_centimetres(self):
        """The CBM is worked from the centimetres printed beside it."""
        services.apply_step(self.order, self.tpl, self.portland, count=1)
        header, body, _ = self.table()
        row = dict(zip(header, body[0]))

        self.assertEqual(
            [row["L (in)"], row["W (in)"], row["H (in)"]], [28, 20, 16]
        )
        self.assertEqual(
            [row["L (cm)"], row["W (cm)"], row["H (cm)"]], [71.12, 50.8, 40.64]
        )
        self.assertEqual(row["CBM"], round(71.12 * 50.8 * 40.64 / 1_000_000, 4))

    def test_the_order_total_counts_every_box_once(self):
        services.apply_step(self.order, self.tpl, self.portland, count=1)
        services.apply_step(self.order, self.tpl, self.austin, count=1)

        header, _, total = self.table()
        self.assertEqual(total[0], "ORDER TOTAL · all 2 boxes")
        self.assertEqual(total[header.index("Qty / Box")], 60)
        self.assertEqual(total[header.index("Total No of Boxes")], 2)

    def test_every_store_reads_as_one_unbroken_ascending_run(self):
        """
        The stored numbers are assigned in step order, which interleaves the
        stores. The sheet runs store by store, so it numbers the rows as it
        writes them — down the whole page, first cell to last.
        """
        self.order.lines.update(quantity=90)
        # Interleaved on purpose: Portland, Austin, Portland again.
        services.apply_step(self.order, self.tpl, self.portland, count=2)
        services.apply_step(self.order, self.tpl, self.austin, count=3)
        services.apply_step(self.order, self.tpl, self.portland, count=1)

        self.assertEqual(
            self.column("Carton Nos"), ["1 – 3", "4 – 6"]
        )

    def test_the_numbering_does_not_restart_in_the_second_store(self):
        """Fifty BOX-001s in one shipment is the first miscount — §10.6."""
        services.apply_step(self.order, self.tpl, self.portland, count=1)
        services.apply_step(self.order, self.tpl, self.austin, count=1)

        self.assertEqual(self.column("Carton Nos"), [1, 2])

    def test_a_store_with_no_cartons_is_left_out(self):
        services.apply_step(self.order, self.tpl, self.portland, count=1)

        self.assertEqual(self.column("Store No"), [118])

    def test_descriptions_print_in_capitals_however_they_were_typed(self):
        DisplayProduct.objects.filter(pk=self.bow.pk).update(
            description="velvet bow", customs_description="100% polyester decorative bow"
        )
        services.apply_step(self.order, self.tpl, self.portland, count=1)

        self.assertEqual(self.column("Customs Description"), ["VELVET BOW"])
        invoice_row = next(
            row for row in self.sheet("invoice").iter_rows(values_only=True)
            if row[1] == "DSP-BOW-12"
        )
        self.assertEqual(invoice_row[3], "100% POLYESTER DECORATIVE BOW")

    def test_the_iec_code_comes_from_settings(self):
        services.apply_step(self.order, self.tpl, self.portland, count=1)
        header = {**settings.EXPORT_DOCUMENT_HEADER, "iec_code": "0999999999"}

        with override_settings(EXPORT_DOCUMENT_HEADER=header):
            sheet = self.sheet()

        self.assertEqual(under(sheet, "Exporter's Ref No"), "IEC No 0999999999")

    def test_the_invoice_opens_with_the_packing_lists_heading(self):
        """The desk reads the two side by side; only the IEC's caption differs."""
        self.order.export_details = {
            "invoice_number": "DKCP-ABC",
            "invoice_date": "2026-07-01",
            "lc_number": "UPL000358085",
            "lc_date": "2021-04-02",
            "other_consignee": "L&J Transportation\nATTN: Stephanie Witmyer",
            "port_of_loading": "ICD DADRI",
        }
        self.order.save()
        services.apply_step(self.order, self.tpl, self.portland, count=1)
        invoice, packing = self.sheet("invoice"), self.sheet()

        self.assertEqual(invoice["A1"].value, "COMMERCIAL INVOICE")
        iec = f"IEC No {settings.EXPORT_DOCUMENT_HEADER['iec_code']}"
        self.assertEqual(under(invoice, "IEC CODE"), iec)
        self.assertEqual(under(packing, "Exporter's Ref No"), iec)
        self.assertEqual(under(invoice, "Invoice No & Date"), "DKCP-ABC Date : 1-07-2026")
        self.assertEqual(under(invoice, "LC No"), "LC NO. UPL000358085 DATED: 02.04.2021")
        for caption in (
            "Exporter", "Invoice No & Date", "LC No", "Consignee",
            "Other Consignee (Shipp To-)", "PORT OF LOADING",
        ):
            self.assertEqual(under(invoice, caption), under(packing, caption), caption)

    def test_the_invoice_heads_its_table_with_the_order(self):
        services.apply_step(self.order, self.tpl, self.portland, count=1)
        values = {cell.value for row in self.sheet("invoice").iter_rows() for cell in row}

        self.assertIn("INVOICE FOR SPLIT", values)

    def test_the_invoice_marks_the_boxes_as_the_packing_list_numbers_them(self):
        services.apply_step(self.order, self.tpl, self.portland, count=1)
        services.apply_step(self.order, self.tpl, self.austin, count=1)

        marks = next(
            [value for value in row if value is not None]
            for row in self.sheet("invoice").iter_rows(values_only=True)
            if row[0] == "MARKS."
        )
        self.assertEqual(marks, ["MARKS.", "1 – 2", "2 BOXES"])

    def test_the_invoice_prints_the_hts_code_not_the_hsn(self):
        """The buyer's customs clear on the US tariff code, not India's HSN."""
        DisplayProduct.objects.filter(pk=self.bow.pk).update(
            hsn_code="9505", hts_code="9505.10.5020"
        )
        services.apply_step(self.order, self.tpl, self.portland, count=1)

        response = self.client.get(
            f"/api/display-orders/{self.order.id}/invoice-data/"
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["lines"][0][1:3], ["DSP-BOW-12", "9505.10.5020"])


class DisplayOrderNumberTests(TestCase):
    def test_numbers_run_sequentially_within_the_year(self):
        merchant = Merchant.objects.create(code="TRN", name="Terrain Home")
        first = DisplayOrder.objects.create(name="One", merchant=merchant)
        second = DisplayOrder.objects.create(name="Two", merchant=merchant)

        self.assertTrue(first.number.startswith("DP-"))
        self.assertEqual(int(second.number[-4:]) - int(first.number[-4:]), 1)


class TemplateCodeTests(TestCase):
    """The code identifies a design; nobody types it."""

    def test_codes_run_sequentially(self):
        first = PackTemplate.objects.create(name="One")
        second = PackTemplate.objects.create(name="Two")

        self.assertEqual((first.code, second.code), ("T-001", "T-002"))

    def test_codes_are_padded_so_they_sort_as_written(self):
        for _ in range(9):
            PackTemplate.objects.create(name="filler")
        tenth = PackTemplate.objects.create(name="Tenth")

        self.assertEqual(tenth.code, "T-010")
        self.assertEqual(
            list(PackTemplate.objects.values_list("code", flat=True))[-2:],
            ["T-009", "T-010"],
        )

    def test_a_code_given_by_hand_is_kept_and_not_counted(self):
        PackTemplate.objects.create(code="LEGACY-7", name="Hand written")
        after = PackTemplate.objects.create(name="Next")

        self.assertEqual(after.code, "T-001")

    def test_an_edit_keeps_the_code_it_was_given(self):
        template = PackTemplate.objects.create(name="One")
        template.name = "Renamed"
        template.save()

        self.assertEqual(template.code, "T-001")


class DisplayProductShapeTests(APITestCase):
    """A display product is either one piece or several parts, fixed at birth."""

    URL = "/api/display-products/"

    def payload(self, **overrides):
        body = {
            "style_no": "DSP-WRT-24",
            "description": '24" Pine Wreath',
            "is_multi_part": True,
            "parts": [
                display_part(name="Frame", product_weight_kg="0.400", sort_order=0),
                display_part(name="Trim", product_weight_kg="0.150", sort_order=1),
            ],
        }
        body.update(overrides)
        return body

    def test_a_multi_part_product_saves_its_parts(self):
        response = self.client.post(self.URL, self.payload(), format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual([part["name"] for part in response.data["parts"]],
                         ["Frame", "Trim"])

    def test_a_multi_part_product_carries_no_figures_of_its_own(self):
        """Its parts are what get weighed and boxed; the whole never is."""
        response = self.client.post(
            self.URL, self.payload(product_weight_kg="0.550"), format="json"
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("product_weight_kg", response.data)

    def test_a_single_piece_product_cannot_have_parts(self):
        response = self.client.post(
            self.URL, self.payload(is_multi_part=False), format="json"
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("parts", response.data)

    def test_two_parts_are_the_minimum(self):
        body = self.payload()
        body["parts"] = body["parts"][:1]
        response = self.client.post(self.URL, body, format="json")

        self.assertEqual(response.status_code, 400)
        self.assertIn("parts", response.data)

    def test_shape_cannot_change_after_creation(self):
        created = self.client.post(self.URL, self.payload(), format="json").data

        response = self.client.patch(
            f"{self.URL}{created['id']}/", {"is_multi_part": False}, format="json"
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("is_multi_part", response.data)

    def test_editing_parts_keeps_their_rows(self):
        """A recreated part would orphan every template item aimed at it."""
        created = self.client.post(self.URL, self.payload(), format="json").data
        ids = [part["id"] for part in created["parts"]]

        response = self.client.patch(
            f"{self.URL}{created['id']}/",
            {
                "parts": [
                    display_part(id=ids[0], name="Wire frame", sort_order=0),
                    display_part(id=ids[1], name="Trim", sort_order=1),
                ]
            },
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual([part["id"] for part in response.data["parts"]], ids)
        self.assertEqual(response.data["parts"][0]["name"], "Wire frame")

    def test_a_part_dropped_from_the_payload_is_deleted(self):
        created = self.client.post(self.URL, self.payload(), format="json").data
        ids = [part["id"] for part in created["parts"]]

        self.client.patch(
            f"{self.URL}{created['id']}/",
            {
                "parts": [
                    display_part(id=ids[0], name="Frame", sort_order=0),
                    display_part(name="Bow", sort_order=1),
                ]
            },
            format="json",
        )

        self.assertEqual(DisplayProductPart.objects.filter(pk=ids[1]).count(), 0)

    CUSTOMS = ("customs_description", "hsn_code", "hts_code")

    def test_a_part_may_leave_its_customs_fields_blank(self):
        body = self.payload()
        for field in self.CUSTOMS:
            del body["parts"][0][field]
            body["parts"][1][field] = ""
        response = self.client.post(self.URL, body, format="json")

        self.assertEqual(response.status_code, 201, response.data)

    def test_a_single_piece_product_may_leave_its_customs_fields_blank(self):
        body = display_product()
        for field in self.CUSTOMS:
            del body[field]
        response = self.client.post(self.URL, body, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual([response.data[field] for field in self.CUSTOMS], ["", "", ""])

    def test_the_hsn_code_is_kept_beside_the_hts_code(self):
        response = self.client.post(self.URL, display_product(), format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(
            (response.data["hsn_code"], response.data["hts_code"]),
            ("9505", "9505.10.5020"),
        )


class PieceCountingTests(TestCase):
    """
    Demand is counted in pieces. A multi-part product is ordered whole but
    packed part by part, and its parts need not share a carton.
    """

    def setUp(self):
        merchant = Merchant.objects.create(code="UO", name="Urban Outfitters")

        self.table = DisplayProduct.objects.create(
            style_no="DSP-TBL-01", description="Display Table", is_multi_part=True
        )
        self.top = DisplayProductPart.objects.create(
            product=self.table, name="Top", product_weight_kg=Decimal("3.2"), sort_order=0
        )
        self.legs = DisplayProductPart.objects.create(
            product=self.table, name="Legs", product_weight_kg=Decimal("2.1"), sort_order=1
        )

        self.store = store()
        self.order = DisplayOrder.objects.create(name="Test", merchant=merchant)
        DisplayOrderLine.objects.create(
            order=self.order,
            store=self.store,
            product=self.table,
            quantity=10,
        )

    def box(self, part, quantity):
        """Put some of one part in a carton of its own."""
        carton = DisplayCarton.objects.create(
            order=self.order,
            store=self.store,
            carton_no=f"BOX-{self.order.cartons.count() + 1:03d}",
            length_in=Decimal("40"),
            width_in=Decimal("20"),
            height_in=Decimal("6"),
            gross_weight_kg=Decimal("5"),
        )
        DisplayCartonContent.objects.create(
            carton=carton, product=self.table, part=part, quantity=quantity
        )

    def test_ordering_a_multi_part_product_asks_for_every_part(self):
        self.assertEqual(
            services.ordered_quantities(self.order),
            Counter(
                {
                    (self.store.id, self.table.id, self.top.id): 10,
                    (self.store.id, self.table.id, self.legs.id): 10,
                }
            ),
        )

    def test_parts_are_counted_apart_so_the_remainder_names_the_short_one(self):
        self.box(self.top, 10)
        self.box(self.legs, 8)

        remaining = services.remaining_quantities(self.order)

        self.assertEqual(remaining, {(self.store.id, self.table.id, self.legs.id): 2})
        self.assertEqual(
            [row["description"] for row in services.remaining_rows(self.order)],
            ["Display Table — Legs"],
        )

    def test_a_product_is_packed_only_once_its_weakest_part_is(self):
        """Ten tops and eight legs is eight tables, not nine."""
        self.box(self.top, 10)
        self.box(self.legs, 8)

        row = services.reconcile(self.order)[0]

        self.assertEqual(row.packed, 8)
        self.assertFalse(row.is_matched)

    def test_every_part_boxed_matches_the_order(self):
        self.box(self.top, 10)
        self.box(self.legs, 10)

        row = services.reconcile(self.order)[0]

        self.assertEqual(row.packed, 10)
        self.assertTrue(row.is_matched)


class TemplateItemPieceTests(APITestCase):
    """
    A template item names a piece. Two parts of one product may sit in two
    different templates, which is the reason the column exists.
    """

    URL = "/api/pack-templates/"

    def setUp(self):
        self.bow = DisplayProduct.objects.create(
            style_no="DSP-BOW-01", description="Velvet Bow",
            product_weight_kg=Decimal("0.1"),
        )
        self.table = DisplayProduct.objects.create(
            style_no="DSP-TBL-01", description="Display Table", is_multi_part=True
        )
        self.top = DisplayProductPart.objects.create(
            product=self.table, name="Top", product_weight_kg=Decimal("3.2"), sort_order=0
        )
        self.legs = DisplayProductPart.objects.create(
            product=self.table, name="Legs", product_weight_kg=Decimal("2.1"), sort_order=1
        )
        self.other = DisplayProductPart.objects.create(
            product=DisplayProduct.objects.create(
                style_no="DSP-OTH-01", description="Other", is_multi_part=True
            ),
            name="Stray",
        )

    def post(self, items, code="TPL-001"):
        return self.client.post(
            self.URL,
            {
                "code": code,
                "name": code,
                "box_length_in": "40",
                "box_width_in": "20",
                "box_height_in": "6",
                "box_weight_kg": "0.8",
                "packing_material_weight_kg": "0.25",
                "is_library": True,
                "items": items,
            },
            format="json",
        )

    def test_a_part_can_share_a_box_with_a_whole_product(self):
        response = self.post(
            [
                {"product": self.table.id, "part": self.top.id, "quantity": 1},
                {"product": self.bow.id, "quantity": 10},
            ]
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(
            [item["part_name"] for item in response.data["items"]], ["Top", ""]
        )

    def test_two_parts_of_one_product_can_go_to_different_templates(self):
        first = self.post(
            [{"product": self.table.id, "part": self.top.id, "quantity": 1}], "TPL-TOP"
        )
        second = self.post(
            [{"product": self.table.id, "part": self.legs.id, "quantity": 1}], "TPL-LEG"
        )

        self.assertEqual(first.status_code, 201, first.data)
        self.assertEqual(second.status_code, 201, second.data)

    def test_a_multi_part_product_must_say_which_part(self):
        response = self.post([{"product": self.table.id, "quantity": 1}])

        self.assertEqual(response.status_code, 400)
        self.assertIn("part", str(response.data))

    def test_a_single_piece_product_takes_no_part(self):
        response = self.post(
            [{"product": self.bow.id, "part": self.top.id, "quantity": 1}]
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("part", str(response.data))

    def test_a_part_of_another_product_is_refused(self):
        response = self.post(
            [{"product": self.table.id, "part": self.other.id, "quantity": 1}]
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("part", str(response.data))

    def test_the_same_piece_twice_is_refused(self):
        response = self.post(
            [
                {"product": self.table.id, "part": self.top.id, "quantity": 1},
                {"product": self.table.id, "part": self.top.id, "quantity": 2},
            ]
        )

        self.assertEqual(response.status_code, 400)

    def test_both_parts_of_one_product_may_share_a_box(self):
        response = self.post(
            [
                {"product": self.table.id, "part": self.top.id, "quantity": 1},
                {"product": self.table.id, "part": self.legs.id, "quantity": 1},
            ]
        )

        self.assertEqual(response.status_code, 201, response.data)

    def test_template_weights_come_from_the_parts(self):
        """Not from the product, which on a multi-part SKU has no weight."""
        response = self.post(
            [
                {"product": self.table.id, "part": self.top.id, "quantity": 1},
                {"product": self.table.id, "part": self.legs.id, "quantity": 1},
            ]
        )

        # 3.2 + 2.1 contents, + 0.25 padding, + 0.8 box
        self.assertEqual(Decimal(str(response.data["net_weight_kg"])), Decimal("5.550"))
        self.assertEqual(Decimal(str(response.data["gross_weight_kg"])), Decimal("6.350"))


class TemplateDeleteTests(APITestCase):
    """A design is deletable until a plan has packed it into cartons."""

    def setUp(self):
        self.merchant = Merchant.objects.create(code="TRQ", name="Tarique")
        self.store = store()
        self.bow = product("DSP-BOW-12", 0.08, (12, 12, 4))
        self.tpl = template("TPL-001", [(self.bow, 10)])

    def test_an_unused_template_is_deleted(self):
        response = self.client.delete(f"/api/pack-templates/{self.tpl.id}/")

        self.assertEqual(response.status_code, 204)
        self.assertFalse(PackTemplate.objects.filter(pk=self.tpl.pk).exists())

    def test_a_packed_template_is_refused_and_names_the_order(self):
        order = DisplayOrder.objects.create(name="Holiday", merchant=self.merchant)
        DisplayOrderLine.objects.create(
            order=order, store=self.store, product=self.bow, quantity=10
        )
        services.apply_step(order, self.tpl, self.store)

        response = self.client.delete(f"/api/pack-templates/{self.tpl.id}/")

        self.assertEqual(response.status_code, 400)
        self.assertIn(order.number, response.data["detail"])
        self.assertTrue(PackTemplate.objects.filter(pk=self.tpl.pk).exists())


class ApplicableTemplateContentsTests(TestCase):
    """Choosing a box design means seeing which pieces it takes."""

    def setUp(self):
        self.merchant = Merchant.objects.create(code="TRQ", name="Tarique")
        self.store = store()
        self.table = DisplayProduct.objects.create(
            style_no="DSP-TBL-01", description="Display Table", is_multi_part=True
        )
        self.top = DisplayProductPart.objects.create(
            product=self.table, name="Top", product_weight_kg=Decimal("3.2"), sort_order=0
        )
        self.legs = DisplayProductPart.objects.create(
            product=self.table, name="Legs", product_weight_kg=Decimal("2.1"), sort_order=1
        )

        self.order = DisplayOrder.objects.create(name="Holiday", merchant=self.merchant)
        # One line per product; the remaining map explodes it into parts.
        DisplayOrderLine.objects.create(
            order=self.order, store=self.store, product=self.table, quantity=20
        )

        self.tpl = PackTemplate.objects.create(code="TPL-001", name="Table box",
                                               is_library=True)
        PackTemplateItem.objects.create(
            template=self.tpl, product=self.table, part=self.top, quantity=2
        )
        PackTemplateItem.objects.create(
            template=self.tpl, product=self.table, part=self.legs, quantity=4
        )

    def test_each_offered_design_carries_its_pieces_and_quantities(self):
        rows = services.applicable_templates(self.order, self.store.id)

        self.assertEqual(len(rows), 1)
        self.assertEqual(
            rows[0]["contents"],
            [
                {"style_no": "DSP-TBL-01", "part_name": "Top", "quantity": 2},
                {"style_no": "DSP-TBL-01", "part_name": "Legs", "quantity": 4},
            ],
        )


PIXEL_GIF = base64.b64decode(
    b"R0lGODlhAQABAIAAAP///wAAACH5BAEAAAAALAAAAAABAAEAAAICRAEAOw=="
)


def upload(name="photo.gif"):
    return SimpleUploadedFile(name, PIXEL_GIF, content_type="image/gif")


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class DisplayProductImageTests(APITestCase):
    """Photos hang off a display product or one of its parts, never both."""

    URL = "/api/display-product-images/"

    def setUp(self):
        self.wreath = DisplayProduct.objects.create(
            style_no="DSP-WRT-24", description='24" Pine Wreath'
        )
        self.arch = DisplayProduct.objects.create(
            style_no="DSP-ARC-84", description='84" Garden Arch', is_multi_part=True
        )
        self.post = DisplayProductPart.objects.create(
            product=self.arch, name="Left post", sort_order=0
        )

    def send(self, **data):
        return self.client.post(self.URL, data, format="multipart")

    def test_a_photo_needs_exactly_one_owner(self):
        self.assertEqual(self.send(image=upload()).status_code, 400)
        self.assertEqual(
            self.send(
                image=upload(), product=self.wreath.pk, part=self.post.pk
            ).status_code,
            400,
        )

    def test_the_first_photo_of_an_owner_becomes_its_main(self):
        first = self.send(image=upload("a.gif"), product=self.wreath.pk)
        second = self.send(image=upload("b.gif"), product=self.wreath.pk)

        self.assertTrue(first.data["is_main"])
        self.assertFalse(second.data["is_main"])

    def test_promoting_a_photo_demotes_the_previous_main(self):
        first = self.send(image=upload("a.gif"), product=self.wreath.pk)
        second = self.send(image=upload("b.gif"), product=self.wreath.pk)

        self.client.patch(
            f"{self.URL}{second.data['id']}/", {"is_main": True}, format="json"
        )

        self.assertFalse(DisplayProductImage.objects.get(pk=first.data["id"]).is_main)
        self.assertTrue(DisplayProductImage.objects.get(pk=second.data["id"]).is_main)

    def test_a_part_carries_its_own_photos(self):
        self.send(image=upload(), part=self.post.pk)

        response = self.client.get(f"/api/display-products/{self.arch.pk}/")

        self.assertEqual(len(response.data["parts"][0]["images"]), 1)
        self.assertEqual(response.data["images"], [])

    def test_a_products_photos_come_back_nested(self):
        self.send(image=upload(), product=self.wreath.pk)

        response = self.client.get(f"/api/display-products/{self.wreath.pk}/")

        self.assertEqual(len(response.data["images"]), 1)
        self.assertIn("/media/display/products/", response.data["images"][0]["image"])


class DisplayStyleNameTests(APITestCase):
    def test_style_name_round_trips(self):
        response = self.client.post(
            "/api/display-products/",
            display_product(
                style_no="DSP-NEW-01",
                style_name="Winter Garland",
                description="Garland",
            ),
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["style_name"], "Winter Garland")

    def test_style_name_is_optional(self):
        response = self.client.post(
            "/api/display-products/",
            display_product(style_no="DSP-NEW-02", description="No name"),
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["style_name"], "")


class DisplayOrderEditTests(APITestCase):
    """
    An order can always be asked for more. Taking away what is already in a
    carton is refused, because reconciliation and the store overview are both
    read off the lines — a packed line that vanishes takes its boxes out of
    every check that would have caught them.
    """

    def setUp(self):
        self.merchant = Merchant.objects.create(code="TRQ", name="Tarique")
        self.store = store()
        self.bow = product("DSP-BOW-12", 0.08, (12, 12, 4))
        self.wreath = product("DSP-WRT-24", 1.4, (24, 24, 6))
        self.tpl = template("TPL-001", [(self.bow, 10)])

        self.order = DisplayOrder.objects.create(name="Holiday", merchant=self.merchant)
        DisplayOrderLine.objects.create(
            order=self.order, store=self.store, product=self.bow, quantity=40
        )

    def line(self, item, quantity):
        return {"store": self.store.id, "product": item.id, "quantity": quantity}

    def put(self, lines):
        return self.client.put(
            f"/api/display-orders/{self.order.id}/",
            {
                "name": self.order.name,
                "merchant": self.merchant.id,
                "buyer_name": "Anthropologie Buying",
                "lines": lines,
            },
            format="json",
        )

    def test_a_product_can_be_added_once_packing_has_started(self):
        services.apply_step(self.order, self.tpl, self.store)

        response = self.put([self.line(self.bow, 40), self.line(self.wreath, 12)])

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(self.order.lines.count(), 2)
        self.assertEqual(self.order.cartons.count(), 4)

    def test_an_untouched_line_keeps_its_row(self):
        before = self.order.lines.get(product=self.bow).id

        self.put([self.line(self.bow, 40), self.line(self.wreath, 12)])

        self.assertEqual(self.order.lines.get(product=self.bow).id, before)

    def test_removing_a_packed_line_is_refused(self):
        services.apply_step(self.order, self.tpl, self.store)

        response = self.put([self.line(self.wreath, 12)])

        self.assertEqual(response.status_code, 400)
        self.assertIn("DSP-BOW-12", str(response.data["lines"]))
        self.assertTrue(self.order.lines.filter(product=self.bow).exists())
        # The whole edit rolled back, so the new line never landed either.
        self.assertFalse(self.order.lines.filter(product=self.wreath).exists())

    def test_reducing_below_what_is_packed_is_refused(self):
        services.apply_step(self.order, self.tpl, self.store)

        response = self.put([self.line(self.bow, 20)])

        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.order.lines.get(product=self.bow).quantity, 40)

    def test_reducing_to_exactly_what_is_packed_is_allowed(self):
        services.apply_step(self.order, self.tpl, self.store, count=2)

        response = self.put([self.line(self.bow, 20)])

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(self.order.lines.get(product=self.bow).quantity, 20)

    def test_an_unpacked_line_is_removed_freely(self):
        response = self.put([self.line(self.wreath, 12)])

        self.assertEqual(response.status_code, 200, response.data)
        self.assertFalse(self.order.lines.filter(product=self.bow).exists())

    def test_a_shipped_order_refuses_edits(self):
        self.order.status = "shipped"
        self.order.save(update_fields=["status"])

        response = self.put([self.line(self.bow, 40)])

        self.assertEqual(response.status_code, 409)

    def test_new_demand_takes_a_packed_order_back_to_packing(self):
        services.apply_step(self.order, self.tpl, self.store)
        self.order.status = "packed"
        self.order.save(update_fields=["status"])

        response = self.put([self.line(self.bow, 40), self.line(self.wreath, 12)])

        self.assertEqual(response.status_code, 200, response.data)
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, "packing")

    def test_the_floor_is_the_largest_piece_count_not_the_smallest(self):
        table = DisplayProduct.objects.create(
            style_no="DSP-TBL-01", description="Table", is_multi_part=True
        )
        top = DisplayProductPart.objects.create(
            product=table, name="Top", product_weight_kg=Decimal("3.2")
        )
        DisplayProductPart.objects.create(
            product=table, name="Legs", product_weight_kg=Decimal("2.1")
        )
        DisplayOrderLine.objects.create(
            order=self.order, store=self.store, product=table, quantity=20
        )
        tops_only = PackTemplate.objects.create(
            code="TPL-TOP", name="Tops", is_library=True
        )
        PackTemplateItem.objects.create(
            template=tops_only, product=table, part=top, quantity=5
        )
        services.apply_step(self.order, tops_only, self.store)

        # Twenty tops are boxed and no legs are. Cutting the line to ten would
        # strand ten of those tops, so twenty is the floor.
        floors = services.packed_floor(self.order)
        self.assertEqual(floors[(self.store.id, table.id)], 20)

        # The plan carries it, so the order form locks on the same number the
        # server enforces. `packed` stays 0 — no whole table is boxed yet.
        row = next(
            r for r in services.reconcile(self.order) if r.product == table.id
        )
        self.assertEqual((row.packed, row.packed_floor), (0, 20))

    def test_a_line_pinned_by_one_part_alone_cannot_be_cut(self):
        table = DisplayProduct.objects.create(
            style_no="DSP-TBL-02", description="Table", is_multi_part=True
        )
        top = DisplayProductPart.objects.create(
            product=table, name="Top", product_weight_kg=Decimal("3.2")
        )
        DisplayProductPart.objects.create(
            product=table, name="Legs", product_weight_kg=Decimal("2.1")
        )
        DisplayOrderLine.objects.create(
            order=self.order, store=self.store, product=table, quantity=20
        )
        tops_only = PackTemplate.objects.create(
            code="TPL-TOP", name="Tops", is_library=True
        )
        PackTemplateItem.objects.create(
            template=tops_only, product=table, part=top, quantity=5
        )
        services.apply_step(self.order, tops_only, self.store)

        response = self.put(
            [self.line(self.bow, 40), {"store": self.store.id, "product": table.id, "quantity": 10}]
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("DSP-TBL-02", str(response.data["lines"]))


class TemplateEditTests(APITestCase):
    """
    Cartons copy their contents when a step is applied, so editing a design
    that is already packed has to rebuild them. A design two orders have
    packed cannot be edited at all — one of them would be rewritten unasked.
    """

    def setUp(self):
        self.merchant = Merchant.objects.create(code="TRQ", name="Tarique")
        self.store = store()
        self.bow = product("DSP-BOW-12", 0.08, (12, 12, 4))
        self.wreath = product("DSP-WRT-24", 1.4, (24, 24, 6))
        self.tpl = template("TPL-001", [(self.bow, 10)])

        self.order = DisplayOrder.objects.create(name="Holiday", merchant=self.merchant)
        DisplayOrderLine.objects.create(
            order=self.order, store=self.store, product=self.bow, quantity=40
        )
        DisplayOrderLine.objects.create(
            order=self.order, store=self.store, product=self.wreath, quantity=8
        )

    def put(self, items, name="Bow box"):
        return self.client.put(
            f"/api/pack-templates/{self.tpl.id}/",
            {"code": "TPL-001", "name": name, "is_library": True, "items": items},
            format="json",
        )

    def test_an_unpacked_design_edits_without_a_replay(self):
        response = self.put([{"product": self.bow.id, "quantity": 5}])

        self.assertEqual(response.status_code, 200, response.data)
        self.assertNotIn("adjustments", response.data)

    def test_adding_a_piece_rebuilds_the_boxes_already_packed(self):
        services.apply_step(self.order, self.tpl, self.store)
        self.assertEqual(self.order.cartons.count(), 4)

        response = self.put(
            [
                {"product": self.bow.id, "quantity": 10},
                {"product": self.wreath.id, "quantity": 2},
            ]
        )

        self.assertEqual(response.status_code, 200, response.data)
        # 40 bows and 8 wreaths both allow four boxes, so the step still stands.
        self.assertEqual(response.data["adjustments"], [])
        self.assertEqual(self.order.cartons.count(), 4)
        self.assertEqual(
            DisplayCartonContent.objects.filter(
                carton__order=self.order, product=self.wreath
            ).count(),
            4,
        )

    def test_a_piece_that_no_longer_fits_cuts_the_step_and_reports_it(self):
        services.apply_step(self.order, self.tpl, self.store)

        response = self.put(
            [
                {"product": self.bow.id, "quantity": 10},
                {"product": self.wreath.id, "quantity": 4},
            ]
        )

        self.assertEqual(response.status_code, 200, response.data)
        # Eight wreaths at four a box is two boxes, not four.
        self.assertEqual(
            response.data["adjustments"],
            [{"sequence": 1, "template_code": "TPL-001", "was": 4, "now": 2}],
        )
        self.assertEqual(self.order.cartons.count(), 2)

    def test_a_rename_leaves_the_plan_alone(self):
        services.apply_step(self.order, self.tpl, self.store)

        response = self.put(
            [{"product": self.bow.id, "quantity": 10}], name="Renamed box"
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertNotIn("adjustments", response.data)
        self.assertEqual(self.order.cartons.count(), 4)

    def test_a_design_two_orders_have_packed_is_refused(self):
        second = DisplayOrder.objects.create(name="Spring", merchant=self.merchant)
        DisplayOrderLine.objects.create(
            order=second, store=self.store, product=self.bow, quantity=20
        )
        services.apply_step(self.order, self.tpl, self.store)
        services.apply_step(second, self.tpl, self.store)

        response = self.put([{"product": self.bow.id, "quantity": 5}])

        self.assertEqual(response.status_code, 400)
        self.assertIn(self.order.number, response.data["detail"])
        self.assertIn(second.number, response.data["detail"])
        self.assertEqual(self.tpl.items.get().quantity, 10)

    def test_a_shipped_order_locks_the_designs_it_packed(self):
        services.apply_step(self.order, self.tpl, self.store)
        self.order.status = "shipped"
        self.order.save(update_fields=["status"])

        response = self.put([{"product": self.bow.id, "quantity": 5}])

        self.assertEqual(response.status_code, 400)
        self.assertIn(self.order.number, response.data["detail"])
