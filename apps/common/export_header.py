"""
The heading block an export document opens with, laid out as the shipping
desk's own: our letterhead, the consignee, the statutory codes and the terms
from `settings.EXPORT_DOCUMENT_HEADER`, and the shipment's own details —
invoice number, LC, ports, vessel — from the order's `export_details`.

Drawn on whatever columns the document already has, so a `Grid` says where
its three bands fall on them.
"""

from dataclasses import dataclass
from datetime import datetime

from django.conf import settings
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RULE = Side(style="thin", color="000000")

YELLOW = PatternFill("solid", fgColor="EFE924")
WHITE = PatternFill("solid", fgColor="FFFFFF")

LABEL = Font(name="Calibri", size=8, bold=True, color="000000")
VALUE = Font(name="Calibri", size=9, bold=False, color="000000")
VALUE_BOLD = Font(name="Calibri", size=9, bold=True, color="000000")
STATUTORY = Font(name="Calibri", size=8.5, bold=True, color="000000")
TITLE = Font(name="Calibri", size=11, bold=True, italic=True, color="000000")

LEFT = Alignment(horizontal="left", vertical="center")
CENTRE = Alignment(horizontal="center", vertical="center")
WRAPPED = Alignment(horizontal="left", vertical="top", wrap_text=True)


@dataclass(frozen=True)
class Grid:
    """
    The last column of each band. Left holds the parties and the transport
    boxes, split in two at `split`; the middle holds the invoice number and
    the other consignee; the right holds the LC and the statutory codes, their
    labels up to `labels`.

    `beside` is where the invoice-number box ends when the reference box fits
    next to it. A middle band one column wide cannot split, so None stacks
    the reference box above the LC box instead.
    """

    left: int
    split: int
    middle: int
    beside: int | None
    labels: int
    last: int


def draw(
    sheet,
    grid: Grid,
    *,
    top: int,
    reference: str,
    title: str,
    export_details: dict,
) -> int:
    """
    Draw the heading from row `top` down and hand back the first free row.
    `reference` captions the box the IEC prints in.
    """
    header = settings.EXPORT_DOCUMENT_HEADER
    static = YELLOW if header["highlight_static"] else None
    entered = WHITE if header["highlight_static"] else None

    def detail(key: str) -> str:
        return str((export_details or {}).get(key) or "").strip()

    left = (1, grid.left)
    carriage = (1, grid.split)
    receipt = (grid.split + 1, grid.left)
    middle = (grid.left + 1, grid.middle)
    right = (grid.middle + 1, grid.last)
    right_labels = (grid.middle + 1, grid.labels)
    right_values = (grid.labels + 1, grid.last)

    # ── Parties and references, five rows ───────────────────────────
    row = top
    block(sheet, row, row + 4, left, "Exporter", "\n".join(header["exporter"]), static)

    iec = header["exporter_ref_no"]
    if grid.beside:
        number = (grid.left + 1, grid.beside)
        block(sheet, row, row + 1, number, "Invoice No & Date", _invoice_no(detail), entered)
        _po(sheet, row + 2, number, detail("po_number"), entered)
        block(sheet, row, row + 2, (grid.beside + 1, grid.middle), reference, iec, static)
        block(sheet, row + 3, row + 4, middle, "Other Reference", detail("other_reference"), entered)
        block(sheet, row, row + 2, right, "LC No", _lc(detail), entered)
        _gstin(sheet, (row + 3, row + 4), right_labels, right_values, header["gstin"], static)
    else:
        block(sheet, row, row + 1, middle, "Invoice No & Date", _invoice_no(detail), entered)
        _po(sheet, row + 2, middle, detail("po_number"), entered)
        block(sheet, row + 3, row + 4, middle, "Other Reference", detail("other_reference"), entered)
        block(sheet, row, row + 1, right, reference, iec, static)
        block(sheet, row + 2, row + 3, right, "LC No", _lc(detail), entered)
        _gstin(sheet, (row + 4, row + 4), right_labels, right_values, header["gstin"], static)

    for line in range(row, row + 5):
        sheet.row_dimensions[line].height = 15

    # ── Consignees, origin and the statutory codes, nine rows ────────
    row = top + 5
    consignee = "\n".join(header["consignee"])
    block(sheet, row, row + 8, left, "Consignee", consignee, static)
    block(
        sheet, row, row + 6, middle, "Other Consignee (Shipp To-)",
        detail("other_consignee"), entered,
    )
    block(
        sheet, row + 7, row + 8, middle, "COUNTRY OF ORIGIN OF GOODS",
        header["country_of_origin"], static, LEFT,
    )

    for offset, (label, value) in enumerate(header["statutory_details"]):
        line = row + offset
        merge(sheet, line, right_labels[0], line, right_labels[1])
        caption = sheet.cell(row=line, column=right_labels[0], value=label)
        caption.font = STATUTORY
        caption.alignment = LEFT
        merge(sheet, line, right_values[0], line, right_values[1])
        figure = sheet.cell(row=line, column=right_values[0], value=value or None)
        figure.font = STATUTORY
        figure.alignment = CENTRE
    box(sheet, row, right[0], row + 6, right[1], static)

    block(
        sheet, row + 7, row + 8, right, "COUNTRY OF FINAL DESTINATION",
        detail("final_destination_country"), entered, LEFT,
    )

    for line in range(row, row + 9):
        sheet.row_dimensions[line].height = 14

    # ── Transport, terms and marks, six rows ─────────────────────────
    row = top + 14
    block(sheet, row, row + 1, carriage, "PRE CARRIAGE BY", header["pre_carriage_by"], static, LEFT)
    block(sheet, row, row + 1, receipt, "PLACE OF RECEIPT", detail("place_of_receipt"), entered, LEFT)
    block(sheet, row + 2, row + 3, carriage, "VESSEL/FLIGHT NO", detail("vessel_flight_no"), entered, LEFT)
    block(sheet, row + 2, row + 3, receipt, "PORT OF LOADING", detail("port_of_loading"), entered, LEFT)
    block(sheet, row + 4, row + 5, carriage, "PORT OF DISCHARGE", detail("port_of_discharge"), entered, LEFT)
    block(sheet, row + 4, row + 5, receipt, "FINAL DESTINATION", detail("final_destination"), entered, LEFT)

    # Settings give the terms every shipment shares; the line and the
    # container change per shipment, so the order's own win where given.
    per_shipment = {
        "SHIPPING LINE": detail("shipping_line"),
        "CONTAINER NO": detail("container_no"),
    }
    for offset, (label, value) in enumerate(header["terms_and_marks"]):
        line = row + offset
        merge(sheet, line, middle[0], line, middle[1])
        caption = sheet.cell(row=line, column=middle[0], value=label)
        caption.font = LABEL
        caption.alignment = LEFT
        merge(sheet, line, right[0], line, right[1])
        entry = sheet.cell(
            row=line, column=right[0], value=per_shipment.get(label) or value or None
        )
        entry.font = VALUE
        entry.alignment = LEFT
        fill = static if value else entered
        if fill:
            for column in range(middle[0], right[1] + 1):
                sheet.cell(row=line, column=column).fill = fill

    merge(sheet, row + 5, middle[0], row + 5, middle[1])
    merge(sheet, row + 5, right[0], row + 5, right[1])
    box(sheet, row, middle[0], row + 5, right[1])

    for line in range(row, row + 6):
        sheet.row_dimensions[line].height = 14

    # ── The document's own title line ────────────────────────────────
    row = top + 20
    merge(sheet, row, 1, row, grid.last)
    heading = sheet.cell(row=row, column=1, value=title)
    heading.font = TITLE
    heading.alignment = LEFT
    box(sheet, row, 1, row, grid.last)
    sheet.row_dimensions[row].height = 24

    return row + 1


def merge(sheet, first_row: int, first_column: int, last_row: int, last_column: int) -> None:
    if (first_row, first_column) != (last_row, last_column):
        sheet.merge_cells(
            start_row=first_row,
            start_column=first_column,
            end_row=last_row,
            end_column=last_column,
        )


def box(sheet, first_row, first_column, last_row, last_column, fill=None) -> None:
    """Rule round the outside of a range only, and fill it when asked."""
    for row in range(first_row, last_row + 1):
        for column in range(first_column, last_column + 1):
            cell = sheet.cell(row=row, column=column)
            cell.border = Border(
                top=RULE if row == first_row else None,
                bottom=RULE if row == last_row else None,
                left=RULE if column == first_column else None,
                right=RULE if column == last_column else None,
            )
            if fill:
                cell.fill = fill


def block(sheet, top, bottom, columns, label, value, fill=None, align=WRAPPED) -> None:
    """A small bold caption over its value, ruled round as one box."""
    first, last = columns
    merge(sheet, top, first, top, last)
    caption = sheet.cell(row=top, column=first, value=label)
    caption.font = LABEL
    caption.alignment = LEFT

    merge(sheet, top + 1, first, bottom, last)
    entry = sheet.cell(row=top + 1, column=first, value=value or None)
    entry.font = VALUE
    entry.alignment = align

    box(sheet, top, first, bottom, last, fill)


def _po(sheet, row, columns, number, fill) -> None:
    """Caption and number share one line, as the desk writes it."""
    first, last = columns
    merge(sheet, row, first, row, last)
    cell = sheet.cell(row=row, column=first, value=f"P.O. No. {number}" if number else "P.O. No.")
    cell.font = LABEL
    cell.alignment = LEFT
    box(sheet, row, first, row, last, fill)


def _gstin(sheet, rows, labels, values, gstin, fill) -> None:
    first, last = rows
    merge(sheet, first, labels[0], last, labels[1])
    caption = sheet.cell(row=first, column=labels[0], value="GSTIN")
    caption.font = LABEL
    caption.alignment = CENTRE
    box(sheet, first, labels[0], last, labels[1], fill)

    merge(sheet, first, values[0], last, values[1])
    number = sheet.cell(row=first, column=values[0], value=gstin)
    number.font = VALUE_BOLD
    number.alignment = CENTRE
    box(sheet, first, values[0], last, values[1], fill)


def _calendar(text: str, pattern: str) -> str:
    """A date from the picker as the desk writes it; anything else as typed."""
    try:
        day = datetime.strptime(text, "%Y-%m-%d")
    except ValueError:
        return text
    return pattern.format(day=day)


def _invoice_no(detail) -> str:
    """`DKCP-ABC Date : 1-07-2026`."""
    number = detail("invoice_number")
    day = _calendar(detail("invoice_date"), "{day.day}-{day.month:02d}-{day.year}")
    if number and day:
        return f"{number} Date : {day}"
    if day:
        return f"Date : {day}"
    return number


def _lc(detail) -> str:
    """`LC NO. UPL000358085 DATED: 02.04.2021`."""
    number = detail("lc_number")
    day = _calendar(detail("lc_date"), "{day.day:02d}.{day.month:02d}.{day.year}")
    if number and not number.upper().startswith("LC NO"):
        number = f"LC NO. {number}"
    if number and day:
        return f"{number} DATED: {day}"
    if day:
        return f"DATED: {day}"
    # Orders saved before the number and the date were two fields.
    return number or detail("lc_number_date")
