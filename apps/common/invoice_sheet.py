"""
The commercial invoice, in the shape the shipping desk already sends.

One row per style across the whole shipment — the packing list says which box
a thing is in, this says what it is worth. The two are read side by side and
their quantities must agree, so both count what is **packed**, never what was
ordered.

Under its COMMERCIAL INVOICE title it opens with the packing list's own
heading (`export_header`), so the pair read alike. The one box they caption
differently holds the IEC: "IEC CODE" here, "Exporter's Ref No" there.
"""

from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from apps.common import export_header

MONEY = Decimal("0.01")
WEIGHT = Decimal("0.001")
ZERO = Decimal("0")

RULE = Side(style="thin", color="000000")
BOX = Border(left=RULE, right=RULE, top=RULE, bottom=RULE)
TITLE = Font(bold=True, size=14)
VALUE = Font(size=10)
HEAD = Font(bold=True, size=9)

# label, width, number format — captioned as the desk captions its own
COLUMNS = [
    ("Serial No", 9, "0"),
    ("Style No", 18, None),
    ("HTS CODE", 14, None),
    ("CUSTOMS DESCRIPTION WITH CONTENTS", 46, None),
    ("Qty in Pcs", 11, "0"),
    ("N.Wt. in Kgs.", 13, "0.000"),
    ("Rate in $", 12, "0.00"),
    ("Amount in US $", 15, "0.00"),
]

SERIAL, STYLE, HTS, DESCRIPTION, QTY, NET, RATE, AMOUNT = range(8)

# Only the description is prose and reads flush left; every code, count and
# figure sits centred in its cell.
CENTRED = Alignment(horizontal="center", vertical="center", wrap_text=True)
FLUSH_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

TITLE_TEXT = "COMMERCIAL INVOICE"

# The middle band is the description column alone, too narrow to split, so
# the IEC box stacks above the LC box — as on the desk's own invoice.
GRID = export_header.Grid(left=3, split=2, middle=4, beside=None, labels=6, last=8)

IGST_NOTE = (
    "Note : IGST @ 18% is paid under refund mechanism and NO CHARGE TO {consignee}. "
    "Supply/service meant for export/supply of sez developer for authorised "
    "operations on payment of IGST"
)
SIGNATORY = "Authorised Signatory"
DECLARATION = [
    "We declare that this Invoice shows the actual Price of goods",
    "described and that all particulars are true and correct.",
    "The above is true to the best of our knowledge",
]


#: The print page still rules these as blank boxes for the desk to fill.
SHIPMENT_FIELDS = [
    ("Invoice No & Date", "P.O. No.", "L/C No"),
    ("Other Reference", "Buyer's Order No", "Country of Final Destination"),
    ("Pre-Carriage By", "Place of Receipt", "Vessel / Flight No"),
    ("Port of Loading", "Port of Discharge", "Final Destination"),
    ("Terms of Delivery", "Payment By", "Shipping Line"),
    ("Container No", "Shipping Mark", "Marks & Numbers"),
]


@dataclass
class Line:
    """One style, summed across every box it ships in."""

    style_no: str
    hts_code: str
    description: str
    quantity: int = 0
    net_weight_kg: Decimal = ZERO
    rate_usd: Decimal | None = None

    @property
    def amount_usd(self) -> Decimal | None:
        if self.rate_usd is None:
            return None
        return (self.rate_usd * self.quantity).quantize(MONEY)

    def row(self, serial: int) -> list:
        return [
            serial,
            self.style_no,
            self.hts_code,
            # In capitals, as the packing list prints the same wording.
            self.description.upper(),
            self.quantity,
            self.net_weight_kg.quantize(WEIGHT),
            self.rate_usd,
            self.amount_usd,
        ]


@dataclass
class Invoice:
    """Every style on the shipment, in style-number order."""

    lines: dict = field(default_factory=dict)

    def add(self, piece, style_no, hts_code, description, quantity, net_weight_kg, rate):
        line = self.lines.get(piece)
        if line is None:
            line = Line(
                style_no=style_no,
                hts_code=hts_code,
                description=description,
                rate_usd=rate,
            )
            self.lines[piece] = line
        # A style shipping in twenty boxes is still one invoice line.
        line.quantity += quantity
        line.net_weight_kg += net_weight_kg or ZERO
        return line

    def ordered(self) -> list[Line]:
        return sorted(self.lines.values(), key=lambda line: line.style_no)

    @property
    def total_quantity(self) -> int:
        return sum(line.quantity for line in self.lines.values())

    @property
    def total_net_weight_kg(self) -> Decimal:
        return sum(
            (line.net_weight_kg for line in self.lines.values()), start=ZERO
        ).quantize(WEIGHT)

    @property
    def total_amount_usd(self) -> Decimal:
        return sum(
            (line.amount_usd or ZERO for line in self.lines.values()), start=ZERO
        ).quantize(MONEY)

    @property
    def is_priced(self) -> bool:
        return any(line.rate_usd is not None for line in self.lines.values())


# ── Sheet ────────────────────────────────────────────────────────────


def set_widths(sheet) -> None:
    for index, (_, width, _) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_title(sheet) -> None:
    last = len(COLUMNS)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    title = sheet.cell(row=1, column=1, value=TITLE_TEXT)
    title.font = TITLE
    title.alignment = Alignment(horizontal="center", vertical="center")
    export_header.box(sheet, 1, 1, 1, last)
    sheet.row_dimensions[1].height = 24


def invoice_title(name: str) -> str:
    """`INVOICE FOR <ORDER>`, the line the desk heads its item table with."""
    name = (name or "").strip().upper()
    return f"INVOICE FOR {name}" if name else "INVOICE"


def write_column_headers(sheet, row: int) -> None:
    for index, (label, _, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=row, column=index, value=label)
        cell.font = HEAD
        cell.border = BOX
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    sheet.row_dimensions[row].height = 30


def write_line(sheet, row: int, values: list) -> None:
    for index, (value, (_, _, number_format)) in enumerate(zip(values, COLUMNS), start=1):
        cell = sheet.cell(row=row, column=index, value=value)
        cell.border = BOX
        cell.font = VALUE
        cell.alignment = FLUSH_LEFT if index - 1 == DESCRIPTION else CENTRED
        if number_format:
            cell.number_format = number_format


def write_marks(sheet, row: int, marks: str, boxes: int) -> int:
    """The boxes the goods travel in, numbered as the packing list numbers them."""
    values = {
        SERIAL: "MARKS.",
        STYLE: marks,
        HTS: f"{boxes} BOX{'ES' if boxes != 1 else ''}",
    }
    for index in range(len(COLUMNS)):
        cell = sheet.cell(row=row, column=index + 1, value=values.get(index))
        cell.border = BOX
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return row + 1


def write_totals(sheet, row: int, invoice: Invoice) -> int:
    # The amount's total stands beside its words further down, where the
    # desk's own invoice puts it, rather than twice over.
    figures = {QTY: invoice.total_quantity, NET: invoice.total_net_weight_kg}

    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=QTY)
    label = sheet.cell(row=row, column=1, value="TOTAL")
    label.font = Font(bold=True)
    label.alignment = Alignment(horizontal="right", vertical="center")

    for index, (_, _, number_format) in enumerate(COLUMNS):
        cell = sheet.cell(row=row, column=index + 1)
        cell.border = BOX
        if index in figures:
            cell.value = figures[index]
            cell.font = Font(bold=True)
            cell.number_format = number_format or "0"
            cell.alignment = CENTRED
    return row + 1


def write_footer(sheet, row: int, invoice: Invoice) -> None:
    last = len(COLUMNS)

    for column in range(1, last + 1):
        sheet.cell(row=row, column=column).border = BOX
    origin = sheet.cell(row=row, column=DESCRIPTION + 1, value=origin_line())
    origin.font = Font(bold=True)
    row += 2

    if invoice.is_priced:
        total = invoice.total_amount_usd
        words = amount_in_words(total)
        _amount_row(sheet, row, f"TOTAL CHARGEABLE AMOUNT IN US DOLLAR — {words}", total)
        # A long amount spelled out runs onto a second line.
        sheet.row_dimensions[row].height = 30
        _amount_row(sheet, row + 1, "Total Amount after tax", total)
        row += 3

    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    note = sheet.cell(row=row, column=1, value=igst_note())
    note.font = Font(bold=True, size=9)
    note.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.row_dimensions[row].height = 30
    row += 2

    declaration = sheet.cell(row=row, column=1, value="Declaration:")
    declaration.font = Font(bold=True, underline="single")
    signatory = sheet.cell(row=row, column=RATE + 1, value=SIGNATORY)
    signatory.font = Font(bold=True)
    row += 1

    for text in DECLARATION:
        sheet.cell(row=row, column=1, value=text).font = Font(size=9)
        row += 1

    exporter = settings.EXPORT_DOCUMENT_HEADER["exporter"][0]
    sheet.cell(row=row, column=RATE + 1, value=f"for {exporter}").font = Font(size=9)


def _amount_row(sheet, row: int, label: str, amount: Decimal) -> None:
    last = len(COLUMNS)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=AMOUNT)
    caption = sheet.cell(row=row, column=1, value=label)
    caption.font = Font(bold=True)
    caption.alignment = Alignment(vertical="center", wrap_text=True)

    figure = sheet.cell(row=row, column=last, value=amount)
    figure.font = Font(bold=True)
    figure.number_format = "0.00"
    figure.alignment = CENTRED

    for column in range(1, last + 1):
        sheet.cell(row=row, column=column).border = BOX


def origin_line() -> str:
    return f"COUNTRY OF ORIGIN OF GOODS - {settings.EXPORT_DOCUMENT_HEADER['country_of_origin']}"


def igst_note() -> str:
    """Names the consignee the refunded IGST is no charge to."""
    return IGST_NOTE.format(consignee=settings.EXPORT_DOCUMENT_HEADER["consignee"][0])


# ── Amount in words ──────────────────────────────────────────────────

_UNITS = (
    "ZERO ONE TWO THREE FOUR FIVE SIX SEVEN EIGHT NINE TEN ELEVEN TWELVE "
    "THIRTEEN FOURTEEN FIFTEEN SIXTEEN SEVENTEEN EIGHTEEN NINETEEN"
).split()
_TENS = (
    "  TWENTY THIRTY FORTY FIFTY SIXTY SEVENTY EIGHTY NINETY".split()
)
_SCALES = [(1_000_000_000, "BILLION"), (1_000_000, "MILLION"), (1_000, "THOUSAND")]


def _under_thousand(number: int) -> str:
    if number < 20:
        return _UNITS[number]
    if number < 100:
        tens, rest = divmod(number, 10)
        word = _TENS[tens - 2]
        return f"{word} {_UNITS[rest]}" if rest else word
    hundreds, rest = divmod(number, 100)
    word = f"{_UNITS[hundreds]} HUNDRED"
    return f"{word} {_under_thousand(rest)}" if rest else word


def _spell(number: int) -> str:
    if number == 0:
        return _UNITS[0]
    parts = []
    for size, name in _SCALES:
        if number >= size:
            count, number = divmod(number, size)
            parts.append(f"{_under_thousand(count)} {name}")
    if number:
        parts.append(_under_thousand(number))
    return " ".join(parts)


def amount_in_words(amount: Decimal) -> str:
    """`TWENTY SEVEN THOUSAND FIVE HUNDRED SEVENTY FIVE AND NINETY FOUR CENTS ONLY`."""
    whole = int(amount)
    cents = int((amount - whole) * 100)
    words = _spell(whole)
    if not cents:
        return f"{words} ONLY"
    return f"{words} AND {_spell(cents)} CENT{'S' if cents != 1 else ''} ONLY"



# ── Document ─────────────────────────────────────────────────────────
#
# The invoice as data. The workbook and the print page render this same
# object, so the sheet emailed to the broker and the PDF signed at the desk
# carry identical figures and identical wording.


@dataclass
class Document:
    invoice: Invoice
    order_title: str = ""
    export_details: dict = field(default_factory=dict)
    # Other Consignee when the export details name none.
    ship_to: list = field(default_factory=list)
    # The shipping marks: the box numbers as the packing list prints them.
    marks: str = ""
    boxes: int = 0


def write_document(sheet, doc: Document) -> None:
    set_widths(sheet)
    write_title(sheet)
    row = export_header.draw(
        sheet,
        GRID,
        top=2,
        reference="IEC CODE",
        title=invoice_title(doc.order_title),
        export_details=doc.export_details,
        ship_to=doc.ship_to,
    )

    write_column_headers(sheet, row)
    row += 1

    for serial, line in enumerate(doc.invoice.ordered(), start=1):
        write_line(sheet, row, line.row(serial))
        row += 1

    row = write_marks(sheet, row, doc.marks, doc.boxes)
    row = write_totals(sheet, row, doc.invoice)
    write_footer(sheet, row, doc.invoice)


def _places(number_format: str | None) -> int | None:
    if not number_format:
        return None
    _, _, fraction = number_format.partition(".")
    return len(fraction)


def _plain(value):
    if isinstance(value, Decimal):
        return float(value)
    return value


def document_json(doc: Document) -> dict:
    header = settings.EXPORT_DOCUMENT_HEADER
    invoice = doc.invoice
    return {
        "title": TITLE_TEXT,
        "exporter": header["exporter"],
        "consignee": header["consignee"],
        "ship_to": export_header.other_consignee(doc.export_details, doc.ship_to),
        "origin": header["country_of_origin"],
        "shipment_fields": [list(group) for group in SHIPMENT_FIELDS],
        "columns": [
            {"label": label, "places": _places(number_format)}
            for label, _, number_format in COLUMNS
        ],
        "lines": [
            [_plain(value) for value in line.row(serial)]
            for serial, line in enumerate(invoice.ordered(), start=1)
        ],
        "total": {
            "quantity": invoice.total_quantity,
            "net_weight_kg": _plain(invoice.total_net_weight_kg),
            "amount_usd": _plain(invoice.total_amount_usd)
            if invoice.is_priced
            else None,
        },
        "is_priced": invoice.is_priced,
        "amount_in_words": amount_in_words(invoice.total_amount_usd)
        if invoice.is_priced
        else "",
        "origin_line": origin_line(),
        "note": igst_note(),
        "declaration": DECLARATION,
        "signatory": SIGNATORY,
        "for_exporter": f"for {header['exporter'][0]}",
    }
