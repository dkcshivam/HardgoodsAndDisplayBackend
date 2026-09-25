"""
The shared machinery behind both packing lists: one row per distinct thing
packed rather than per carton, and the spreadsheet furniture around it.

Hardgoods and Display keep separate models and separate documents, but a
carton is a carton — the grouping, the carton-range notation and the column
layout are the same question in both, and a shipping desk reading one after
the other should not have to learn two conventions.

Everything here is duck-typed against a carton: `carton_no`, `contents`,
`gross_weight_kg`, `length_in`/`width_in`/`height_in`, `cbm`.
"""

import re
from dataclasses import dataclass, field
from decimal import Decimal, ROUND_HALF_UP

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.common import export_header

# label, width, number format
#
# The order and the wording follow the packing list the shipping desk already
# sends out by hand: what is in the box, then its three weights, then the box
# itself. Every figure describes **one** box, however many boxes the row
# stands for — the footer carries the shipment.
COLUMNS = [
    ("Carton Nos", 20, None),
    ("Cartons", 8, "0"),
    ("Style No", 18, None),
    ("Customs Description", 34, None),
    ("Qty / Box", 9, "0"),
    ("Units", 8, None),
    ("NNW (kg)", 11, "0.000"),
    ("N.W. (kg)", 11, "0.000"),
    ("G.W. (kg)", 11, "0.000"),
    ("L (cm)", 9, "0.00"),
    ("W (cm)", 9, "0.00"),
    ("H (cm)", 9, "0.00"),
    ("CBM", 10, "0.0000"),
]

# What the footer fills in. Not a column sum: a row standing for twelve
# identical boxes prints one box's figures and has to count as twelve, so
# these are worked out from the grouping instead. Dimensions total to nothing
# meaningful and are left blank.
TOTALLED = {
    "Cartons",
    "Qty / Box",
    "NNW (kg)",
    "N.W. (kg)",
    "G.W. (kg)",
    "CBM",
}

CM_PER_INCH = Decimal("2.54")
CM_PLACES = Decimal("0.01")
CUBIC_CM_PER_CBM = Decimal("1000000")
CBM_PLACES = Decimal("0.0001")


def cm(inches) -> Decimal | None:
    """
    Inches as the sheet prints them. Two places, because an inch is exactly
    2.54 cm — a whole-inch box converts without any rounding at all, and the
    CBM below then matches the one the app worked out in inches.
    """
    if inches is None:
        return None
    return (Decimal(inches) * CM_PER_INCH).quantize(
        CM_PLACES, rounding=ROUND_HALF_UP
    )


def cbm_from_cm(length_in, width_in, height_in) -> Decimal | None:
    """
    Worked from the centimetres actually printed rather than from the inches
    behind them, so a broker who multiplies the three numbers on the page
    arrives at the fourth.
    """
    sides = [cm(value) for value in (length_in, width_in, height_in)]
    if any(side is None for side in sides):
        return None
    length, width, height = sides
    return (length * width * height / CUBIC_CM_PER_CBM).quantize(CBM_PLACES)

@dataclass(frozen=True)
class Layout:
    """
    Which columns a sheet prints, and which of them carry a total.

    `Group.rows()` always builds the full canonical row above; a layout that
    prints something else — Display drops the carton count, because its rows
    are already blocked under the store they ship to — takes the columns it
    wants, so the grouping never has to know.
    """

    columns: list
    totalled: set
    header: export_header.Grid


HEADER_FILL = PatternFill("solid", fgColor="EEF1FF")
GROUP_FILL = PatternFill("solid", fgColor="F5F6FA")
RULE = Side(style="thin", color="D8DCE6")
BORDER = Border(left=RULE, right=RULE, top=RULE, bottom=RULE)

# A box's rows read as one thing, so the rule that opens a box is heavier
# than the ones between its own styles.
EDGE = Side(style="medium", color="8B93A3")

#: What a sheet prints unless it says otherwise — Hardgoods uses this as is.
BASE = Layout(
    COLUMNS,
    TOTALLED,
    header=export_header.Grid(left=3, split=1, middle=7, beside=5, labels=9, last=13),
)


# ── Grouping ─────────────────────────────────────────────────────────


@dataclass
class Group:
    """Cartons that are the same box holding the same thing, so one row."""

    contents: list
    cartons: list = field(default_factory=list)
    # Set where the document numbers its own rows; empty prints the numbers
    # stored on the cartons.
    label: str | int = ""

    def rows(self) -> list[list]:
        count = len(self.cartons)
        sample = self.cartons[0]
        label = self.label or carton_range(self.cartons)

        if not self.contents:
            return [empty_carton_row(sample, label)]

        built = []
        for position, content in enumerate(self.contents):
            # Only the first row of a box carries the box's own figures. A
            # carton holding six different things would otherwise state its
            # weight six times, and read as six boxes.
            opens = position == 0
            built.append(
                [
                    label if opens else "",
                    count if opens else None,
                    content.product.style_no,
                    describe(content),
                    content.quantity,
                    unit_label(content),
                    nnw(content),
                    sample.net_weight_kg if opens else None,
                    sample.gross_weight_kg if opens else None,
                    cm(sample.length_in) if opens else None,
                    cm(sample.width_in) if opens else None,
                    cm(sample.height_in) if opens else None,
                    cbm_from_cm(
                        sample.length_in, sample.width_in, sample.height_in
                    )
                    if opens
                    else None,
                ]
            )
        return built


def group_cartons(cartons: list) -> list[Group]:
    """
    Cartons collapse into one row when the box and its single content match
    on every printed field. A carton holding more than one thing stands
    alone — merging it would need every other carton to hold the same mix.
    """
    groups: dict[tuple, Group] = {}
    ordered: list[Group] = []

    for carton in cartons:
        contents = list(carton.contents.all())
        key = _key(carton, contents)

        group = groups.get(key) if key else None
        if group is None:
            group = Group(contents=contents)
            ordered.append(group)
            if key:
                groups[key] = group

        group.cartons.append(carton)

    return ordered


def _key(carton, contents: list) -> tuple | None:
    """None for anything that must keep a row of its own."""
    if len(contents) != 1:
        return None

    content = contents[0]
    return (
        content.product_id,
        content.part_id,
        content.description,
        content.quantity,
        content.unit,
        content.net_weight_kg,
        carton.gross_weight_kg,
        carton.length_in,
        carton.width_in,
        carton.height_in,
    )


TRAILING_NUMBER = re.compile(r"(\d+)$")


def carton_range(cartons: list) -> str:
    """
    `BOX-001 – BOX-012` for a run, and each run named when a hand-edited
    plan leaves gaps: `BOX-001 – BOX-003, BOX-007`.

    Every number a row covers is either printed or inside a printed run.
    A shorthand for the gaps would be shorter, but a range that has to be
    decoded is one somebody miscounts, and this document is read by people
    who will not ask.
    """
    labels = [(carton.carton_no or "").strip() for carton in cartons]
    if len(labels) == 1:
        return labels[0]

    numbered = [TRAILING_NUMBER.search(label) for label in labels]
    if not all(numbered):
        return ", ".join(labels)

    pairs = sorted(zip((int(match.group(1)) for match in numbered), labels))
    numbers = [number for number, _ in pairs]

    runs, start = [], 0
    for index in range(1, len(pairs) + 1):
        if index == len(pairs) or numbers[index] != numbers[index - 1] + 1:
            runs.append((start, index - 1))
            start = index

    return ", ".join(
        pairs[a][1] if a == b else f"{pairs[a][1]} – {pairs[b][1]}" for a, b in runs
    )


def number_groups(groups: list[Group], start: int = 1) -> int:
    """
    Number the rows in the order the sheet writes them, and hand back the
    next free number.

    `carton_no` runs in step order, and a sheet blocked by store reads those
    out of sequence — store 118 printing `BOX-001 – BOX-002, BOX-006` above a
    store whose run starts at 003. Numbering as we write gives every row one
    unbroken run and the page one ascending column. `start` threads across the
    blocks so the shipment keeps a single sequence (§10.6).

    Plain numbers, as the shipping desk writes them: a lone box is the number
    itself, so Excel stores it as one, and a run reads `10 – 14`.
    """
    for group in groups:
        last = start + len(group.cartons) - 1
        group.label = start if start == last else f"{start} – {last}"
        start = last + 1
    return start


def times(value: Decimal | None, count: int) -> Decimal | None:
    return None if value is None else value * count


def unit_label(content) -> str:
    """`pcs` on the record, `PCS` on a customs document."""
    return (content.unit or "").upper()


WEIGHT_PLACES = Decimal("0.001")


def nnw(content) -> Decimal | None:
    """
    The goods alone — unit weight times how many are in the box, with no
    packing material and no carton.

    `content.net_weight_kg` cannot stand in for this: Hardgoods folds the
    padding into it and Display does not, so the column would mean two
    different things on two documents the same desk reads.
    """
    piece = content.part if content.part_id else content.product
    weight = getattr(piece, "product_weight_kg", None)
    if weight is None:
        return None
    return (Decimal(weight) * (content.quantity or 0)).quantize(WEIGHT_PLACES)


@dataclass
class Totals:
    """
    What the shipment comes to, gathered as the rows are written.

    It cannot be a column sum. Every printed figure is one box's, and a row
    may stand for twelve of them, so each group is counted by its run.
    """

    cartons: int = 0
    quantity: int = 0
    nnw_kg: Decimal = Decimal("0")
    net_kg: Decimal = Decimal("0")
    gross_kg: Decimal = Decimal("0")
    cbm: Decimal = Decimal("0")

    def add(self, group: "Group") -> None:
        count = len(group.cartons)
        sample = group.cartons[0]
        self.cartons += count

        for content in group.contents:
            self.quantity += (content.quantity or 0) * count
            self.nnw_kg += (nnw(content) or Decimal("0")) * count

        self.net_kg += (sample.net_weight_kg or Decimal("0")) * count
        self.gross_kg += (sample.gross_weight_kg or Decimal("0")) * count
        volume = cbm_from_cm(sample.length_in, sample.width_in, sample.height_in)
        self.cbm += (volume or Decimal("0")) * count

    def merge(self, other: "Totals") -> None:
        self.cartons += other.cartons
        self.quantity += other.quantity
        self.nnw_kg += other.nnw_kg
        self.net_kg += other.net_kg
        self.gross_kg += other.gross_kg
        self.cbm += other.cbm

    def by_column(self) -> dict:
        return {
            "Cartons": self.cartons,
            "Total No of Boxes": self.cartons,
            "Qty / Box": self.quantity,
            "NNW (kg)": self.nnw_kg,
            "N.W. (kg)": self.net_kg,
            "G.W. (kg)": self.gross_kg,
            "CBM": self.cbm,
        }


def describe(content) -> str:
    """
    The customs wording where there is any, and the nearest thing to it where
    there is not. A blank cell here is a document a broker cannot clear, and
    a Display product often carries only a style name.

    Always in capitals, however the catalogue typed it, so a column mixing
    customs wording with style names still reads as one hand.
    """
    if content.description:
        wording = content.description
    elif content.part_id:
        part = content.part
        wording = part.customs_description or f"{_names(content.product)} — {part.name}"
    else:
        wording = _names(content.product)
    return wording.upper()


def _names(product) -> str:
    return (
        product.customs_description
        or getattr(product, "style_name", "")
        or product.description
    )


def empty_carton_row(carton, label: str | None = None) -> list:
    return [
        label if label is not None else carton_range([carton]),
        1,
        "",
        "(empty carton)",
        None,
        "",
        None,
        carton.net_weight_kg,
        carton.gross_weight_kg,
        cm(carton.length_in),
        cm(carton.width_in),
        cm(carton.height_in),
        cbm_from_cm(carton.length_in, carton.width_in, carton.height_in),
    ]


# ── Sheet ────────────────────────────────────────────────────────────


def set_widths(sheet, layout: Layout = None) -> None:
    layout = layout or BASE
    for index, (_, width, _) in enumerate(layout.columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_facts(sheet, row: int, facts: list[tuple[str, str]]) -> int:
    for label, value in facts:
        if not value:
            continue
        sheet.cell(row=row, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=row, column=2, value=value)
        row += 1
    return row


def write_column_headers(sheet, row: int, layout: Layout = None) -> None:
    layout = layout or BASE
    for index, (label, _, _) in enumerate(layout.columns, start=1):
        cell = sheet.cell(row=row, column=index, value=label)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )


def write_row(
    sheet, row: int, values: list, layout: Layout = None, opens: bool = False
) -> None:
    """`opens` draws the heavier rule that separates one box from the last."""
    layout = layout or BASE
    border = (
        Border(left=RULE, right=RULE, top=EDGE, bottom=RULE) if opens else BORDER
    )
    for index, (value, (_, _, number_format)) in enumerate(
        zip(values, layout.columns), start=1
    ):
        cell = sheet.cell(row=row, column=index, value=value)
        cell.border = border
        if number_format:
            cell.number_format = number_format
            cell.alignment = Alignment(horizontal="right")


def write_banner(sheet, row: int, text: str, layout: Layout = None) -> None:
    """A full-width label opening a block — the store a run of rows is for."""
    layout = layout or BASE
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = Font(bold=True)
    cell.fill = GROUP_FILL
    for index in range(1, len(layout.columns) + 1):
        sheet.cell(row=row, column=index).fill = GROUP_FILL


def write_totals(
    sheet,
    row: int,
    totals: Totals,
    label: str,
    layout: Layout = None,
) -> None:
    """
    Values, not `=SUM()` formulas. The columns above hold one box's figures
    and a row may stand for twelve, so nothing on the page adds up to what
    ships — `Totals` counts the boxes and these are what it reached.
    """
    layout = layout or BASE
    if totals.cartons == 0:
        return

    heading = sheet.cell(row=row, column=1, value=label)
    heading.font = Font(bold=True)

    figures = totals.by_column()
    for index, (label_text, _, number_format) in enumerate(layout.columns, start=1):
        if label_text not in layout.totalled:
            continue
        cell = sheet.cell(row=row, column=index, value=figures.get(label_text))
        cell.font = Font(bold=True)
        cell.border = Border(top=Side(style="double", color="9AA2B1"))
        cell.number_format = number_format or "0"
        cell.alignment = Alignment(horizontal="right")


# ── Document ─────────────────────────────────────────────────────────
#
# What the packing list says, with nothing about how it is drawn. The
# workbook and the print page both render this, so the .xlsx a broker files
# and the PDF the desk signs cannot drift apart.


@dataclass
class Block:
    """A run of rows under one banner — for Display, one store."""

    rows: list = field(default_factory=list)  # (values, opens)
    banner: str = ""
    subtotal_label: str = ""
    subtotal: "Totals | None" = None


@dataclass
class Document:
    title: str
    facts: list
    blocks: list
    total: Totals
    total_label: str
    layout: Layout = BASE
    order_title: str = ""
    export_details: dict = field(default_factory=dict)
    # Other Consignee when the export details name none.
    ship_to: list = field(default_factory=list)


def packing_title(name: str) -> str:
    """`PACKING LIST FOR <ORDER> ITEMS`, as the desk heads its own."""
    name = (name or "").strip().upper()
    if not name or name.startswith("PACKING LIST"):
        return name or "PACKING LIST"
    return f"PACKING LIST FOR {name}" + ("" if name.endswith("ITEMS") else " ITEMS")


def write_document(sheet, doc: Document) -> None:
    set_widths(sheet, doc.layout)

    header_row = export_header.draw(
        sheet,
        doc.layout.header,
        top=1,
        reference="Exporter's Ref No",
        title=packing_title(doc.order_title or doc.title),
        export_details=doc.export_details,
        ship_to=doc.ship_to,
    )
    write_column_headers(sheet, header_row, doc.layout)
    row = header_row + 1

    for block in doc.blocks:
        if block.banner:
            write_banner(sheet, row, block.banner, doc.layout)
            row += 1
        for values, opens in block.rows:
            write_row(sheet, row, values, doc.layout, opens=opens)
            row += 1
        if block.subtotal is not None:
            write_totals(
                sheet, row, block.subtotal, block.subtotal_label, doc.layout
            )
            row += 2

    write_totals(sheet, row, doc.total, doc.total_label, doc.layout)
    # Note: freeze_panes intentionally omitted so header rows move with the rest of the sheet




def places(number_format: str | None) -> int | None:
    """Decimal places from an Excel format, so the page can round the same."""
    if not number_format:
        return None
    _, _, fraction = number_format.partition(".")
    return len(fraction)


def plain(value):
    """Decimals go out as numbers — the API does not stringify them."""
    if isinstance(value, Decimal):
        return float(value)
    return value


def _totals_row(totals: "Totals | None", layout: Layout) -> list | None:
    if totals is None or totals.cartons == 0:
        return None
    figures = totals.by_column()
    return [
        plain(figures.get(label)) if label in layout.totalled else None
        for label, _, _ in layout.columns
    ]


def document_json(doc: Document) -> dict:
    return {
        "title": doc.title,
        "facts": [
            {"label": label, "value": value} for label, value in doc.facts if value
        ],
        "columns": [
            {"label": label, "places": places(number_format)}
            for label, _, number_format in doc.layout.columns
        ],
        "blocks": [
            {
                "banner": block.banner,
                "rows": [
                    {"opens": opens, "values": [plain(value) for value in values]}
                    for values, opens in block.rows
                ],
                "subtotal_label": block.subtotal_label,
                "subtotal": _totals_row(block.subtotal, doc.layout),
            }
            for block in doc.blocks
        ],
        "total_label": doc.total_label,
        "total": _totals_row(doc.total, doc.layout),
    }
